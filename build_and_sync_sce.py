import os
import sys
import json
import time
import urllib.request
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

records = [
    {
        "序号": 1,
        "姓名": "张廷佺",
        "职称": "教授",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "张廷佺，博士研究生学历，上海外国语大学继续教育学院院长、教授，兼任教育部教育质量评估中心评估专家、上海翻译家协会理事。主要研究方向为英语语言文学、翻译及高校教学管理。主持或参与国家社科基金重点项目多项，主译英美文学作品逾300万字，任多部规划教材总主编，曾获上海市成人高校“十佳院长”及上外教学成果一等奖。",
        "主页": "https://sce.shisu.edu.cn/f8/58/c6830a194648/page.htm"
    },
    {
        "序号": 2,
        "姓名": "鲍晓英",
        "职称": "教授、博士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "鲍晓英，博士研究生学历，上海外国语大学教授、博士生导师，曾任希腊亚里士多德大学孔子学院中方院长、美国蒙特雷国际研究学院访问学者。主要研究方向为英语语言文学与翻译学。入选上海市浦江人才计划，主持教育部及上海市哲社项目5项，在《中国翻译》《外国语》等核心刊物发表论文近30篇，出版专著2部及Springer学术译著2部。",
        "主页": "https://sce.shisu.edu.cn/9c/36/c6830a105526/page.htm"
    },
    {
        "序号": 3,
        "姓名": "罗杏焕",
        "职称": "副教授、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "罗杏焕，硕士研究生学历，上海外国语大学继续教育学院副教授、硕士生导师，现任院教学指导委员会主任委员、院学术委员会副主任委员。主要研究领域为英语语言学、英语教学及成人教育。长期主讲专业阅读与基础综合英语课程，独立或合作出版教材9部、专著2部，发表学术论文10余篇，曾获全国成人教育优秀奖及上海外国语大学育才奖。",
        "主页": "https://sce.shisu.edu.cn/9c/40/c6830a105536/page.htm"
    },
    {
        "序号": 4,
        "姓名": "孙黎",
        "职称": "副教授、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "孙黎，博士研究生学历，上海外国语大学继续教育学院副教授、硕士生导师，院学术委员会委员，曾赴英国利兹大学进修访学。主要研究方向为英美文学与翻译。主讲英美文学、口笔译等课程，获评上海外国语大学“教学十佳”教师。深度参与精品课程建设，参编《文体翻译学》《剑桥美国文学史》等学术专著与专业词典，主编《综合英语教程》等多部重点教材。",
        "主页": "https://sce.shisu.edu.cn/9c/3a/c6830a105530/page.htm"
    },
    {
        "序号": 5,
        "姓名": "黄皓",
        "职称": "副教授、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "黄皓，硕士研究生学历，上海外国语大学继续教育学院副教授、硕士生导师，曾任秘鲁天主教大学孔子学院中方院长、中组部第八批援疆干部。主要研究方向为应用语言学、翻译理论及继续教育外语教学。发表学术论文27篇，编著、翻译专业作品26部（含教材10部、译作9部）。多次荣获上海外国语大学校级教学成果二等奖、优秀教材奖及上海图书奖。",
        "主页": "https://sce.shisu.edu.cn/9c/34/c6830a105524/page.htm"
    },
    {
        "序号": 6,
        "姓名": "赵蓉",
        "职称": "副教授、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "赵蓉，博士研究生学历，上海外国语大学继续教育学院副教授、硕士生导师，研究生教研室主任，校二级学术骨干。主要研究方向为外语教学与电化教育。主持并参与教育部人文社科项目、国家语委“十三五”重点项目及上海市教科规划项目多项。在《电化教育研究》《外语电化教学》等CSSCI核心期刊发表论文多篇，参编口译系列规划教材多部。",
        "主页": "https://sce.shisu.edu.cn/9c/32/c6830a105522/page.htm"
    },
    {
        "序号": 7,
        "姓名": "张昕宇",
        "职称": "副教授、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "张昕宇，博士研究生学历，上海外国语大学继续教育学院副教授、硕士生导师。主要研究方向为日语语言文学及日本近现代文学。在村上春树文学研究领域造诣深厚，在中文核心及海外学术期刊发表论文多篇。参编《日语听力》《新日语教程》等国家级规划教材，主持完成上海市终身教育学分银行等重点教改课题，曾获上海外国语大学优秀教学成果奖。",
        "主页": "https://sce.shisu.edu.cn/9c/31/c6830a105521/page.htm"
    },
    {
        "序号": 8,
        "姓名": "徐丹",
        "职称": "副教授、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "徐丹，硕士研究生学历，上海外国语大学继续教育学院副教授、硕士生导师，校二级学术骨干、上海市中高级口译考官。主要研究方向为成人外语教育及国别研究。在《中国高教研究》《教育发展研究》等CSSCI期刊发表论文多篇，参编《综合英语教程》等多部规划教材与词典，曾获上海市高等学历继续教育优秀教材一等奖及校教学科研成果奖。",
        "主页": "https://sce.shisu.edu.cn/9c/3e/c6830a105534/page.htm"
    },
    {
        "序号": 9,
        "姓名": "彭梅",
        "职称": "副教授、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "彭梅，博士研究生学历，上海外国语大学继续教育学院副教授、硕士生导师，院学术委员会委员，曾任美国匹兹堡大学访问学者、孔子学院公派教师，获评上外“教学之星”。主要研究方向为外语教学理论与实践。在国内外核心刊物发表论文23篇（含SSCI及CSSCI），出版专著1部，主持或参与省部级与厅级课题6项，曾获教育部基础教育课程改革教学成果二等奖。",
        "主页": "https://sce.shisu.edu.cn/9c/3d/c6830a105533/page.htm"
    },
    {
        "序号": 10,
        "姓名": "徐佳",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "徐佳，博士研究生学历，上海外国语大学继续教育学院讲师，教工第一党支部书记、院教学指导委员会委员，曾赴澳大利亚阿德莱德大学访学。主要研究方向为语言学与外语教学。出版学术专著1部，参编《语言学与应用语言学百科全书》（获上海市哲社优秀成果二等奖）及规划教材多部，主持校级科研及精品课程项目多项，获评“上海市成人高校优秀教师”。",
        "主页": "https://sce.shisu.edu.cn/9c/2a/c6830a105514/page.htm"
    },
    {
        "序号": 11,
        "姓名": "曹磊",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "曹磊，博士研究生学历，上海外国语大学继续教育学院讲师，院教学指导委员会委员。主要研究方向为语言学、翻译学及外语教学。出版学术专著《翻译的修辞符号视角研究》，参编《综合英语教程》《高级口译教程》《高级听力教程》等国家级与行业知名英语教材十余部。长期深耕英语综合与口笔译教学，荣获“2023-2024年度上海市成人高校优秀教师”称号。",
        "主页": "https://sce.shisu.edu.cn/9c/2e/c6830a105518/page.htm"
    },
    {
        "序号": 12,
        "姓名": "周艳",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "周艳，硕士研究生学历，上海外国语大学继续教育学院讲师，院教学指导委员会委员。主要研究方向为商务英语与成人外语教育。长期主讲商务英语、高级听说及中高级口译等核心课程。作为骨干编者参编《综合英语教程》《基础口译听力教程》等多部重点教材与大型辞书，曾获“上海市成人高校优秀教师”称号及校教育奖励基金教学科研奖。",
        "主页": "https://sce.shisu.edu.cn/9c/2f/c6830a105519/page.htm"
    },
    {
        "序号": 13,
        "姓名": "贾丹",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "贾丹，博士研究生学历，上海外国语大学继续教育学院讲师，院教学指导委员会委员。主要研究领域为二语习得与英语教学法。参与国家“211工程”重点建设项目及《语言学与应用语言学百科全书》编纂，该成果荣获上海市哲社优秀成果二等奖；参编《中级口译教程》《高级口译教程》等精品教材，荣获校级教学成果二等奖及“上海市成人高校优秀教师”称号。",
        "主页": "https://sce.shisu.edu.cn/9c/35/c6830a105525/page.htm"
    },
    {
        "序号": 14,
        "姓名": "葛明霞",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "葛明霞，硕士研究生学历，上海外国语大学继续教育学院讲师，现任外语教研室副主任。主要研究方向为英语语言文学与写作教学。在《现代外语教学》《外语教学与研究》等领域发表多篇成人英语写作与语用学教研论文，参与编写《综合英语教程》等重点教材。长期深耕英语写作与精读一线教学，曾获校教育奖励基金二等奖及“上海市成人高校优秀教师”称号。",
        "主页": "https://sce.shisu.edu.cn/9c/2d/c6830a105517/page.htm"
    },
    {
        "序号": 15,
        "姓名": "蒋晓艳",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "蒋晓艳，博士研究生学历，上海外国语大学继续教育学院讲师，曾赴加拿大卡普兰诺大学进修。主要研究方向为外语教学理论与实践、外语教师专业发展。出版学术专著1部，发表学术论文数篇，参编《高级口译听力》等多部规划教材。长期从事成人高等教育综合英语与雅思听说教学，深受学生好评，曾荣获上海外国语大学教育奖励基金教学科研奖。",
        "主页": "https://sce.shisu.edu.cn/9c/3b/c6830a105531/page.htm"
    },
    {
        "序号": 16,
        "姓名": "袁捷",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "袁捷，硕士研究生学历，上海外国语大学继续教育学院讲师。主要研究方向为应用语言学与英语教学法。长期从事英语专业阅读与听说课程的一线教学工作，因材施教、治学严谨。参编多部高校英语专业教材，教学成效显著，所指导和培养的学生在全国英语专业四级（TEM-4）和专业八级（TEM-8）统测中屡创优异成绩，深得广大师生好评。",
        "主页": "https://sce.shisu.edu.cn/9c/2c/c6830a105516/page.htm"
    },
    {
        "序号": 17,
        "姓名": "俞丹",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "俞丹，硕士研究生学历，上海外国语大学继续教育学院讲师。主要研究方向为外国语言学与应用语言学。主讲综合英语、高级英语写作、英语视听说等专业核心课程。主持或参与校级教学科研团队培育计划、院级科研管理项目及一流课程建设项目多项，在学术刊物发表《英汉称呼语比较》《翻转课堂模式在成人英语教学中的应用》等教研论文多篇。",
        "主页": "https://sce.shisu.edu.cn/37/e7/c6830a145383/page.htm"
    },
    {
        "序号": 18,
        "姓名": "王峥",
        "职称": "讲师、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "王峥，博士研究生学历，上海外国语大学继续教育学院讲师、硕士生导师，院教学指导委员会秘书。主要研究方向为二语习得与在线英语学习心理。出版专著1部、译著3部，参编教材与词典5部，主持科研项目5项。在CSSCI及EI检索期刊发表学术论文多篇，荣膺全国成人继续教育优秀科研成果一等奖，两次获评“上海市成人教育优秀教师”。",
        "主页": "https://sce.shisu.edu.cn/9c/2b/c6830a105515/page.htm"
    },
    {
        "序号": 19,
        "姓名": "施莉莉",
        "职称": "讲师、硕士生导师",
        "邮箱": "",
        "学校": "上海外国语大学",
        "学院": "继续教育学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "施莉莉，博士研究生学历，上海外国语大学继续教育学院讲师、硕士生导师，兼任上海终身教育研究院兼职研究员。主要研究方向为成人教育、高等教育及区域国别教育（日本）。主持教育部人文社会科学研究青年基金项目，在国内外SSCI、CSSCI等权威核心期刊及国际高水平会议发表论文十余篇，出版学术译著1部，承担本硕多门日语与教育学核心课程。",
        "主页": "https://sce.shisu.edu.cn/1e/b0/c6830a138928/page.htm"
    }
]

columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']

with_email = [r for r in records if r['邮箱'].strip()]
without_email = [r for r in records if not r['邮箱'].strip()]

# Dual-sheet validation
assert len(records) == 19
assert len(with_email) + len(without_email) == 19
print(f"Dual-sheet validation passed: Total={len(records)}, WithEmail={len(with_email)}, WithoutEmail={len(without_email)}")

sheet1_name = f"有邮箱学者 (共 {len(with_email)} 位)"
sheet2_name = f"无邮箱学者 (共 {len(without_email)} 位)"

timestamp = time.strftime('%Y%m%d_%H%M%S')
filename = f"上海外国语大学_继续教育学院_{timestamp}.xlsx"
output_dir = "server/outputs"
os.makedirs(output_dir, exist_ok=True)
filepath = os.path.join(output_dir, filename)

wb = openpyxl.Workbook()
# Sheet 1
ws1 = wb.active
ws1.title = sheet1_name

# Sheet 2
ws2 = wb.create_sheet(title=sheet2_name)

header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
data_font = Font(name="Microsoft YaHei", size=10, color="334155")
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

col_widths = {
    '序号': 8,
    '姓名': 12,
    '职称': 18,
    '邮箱': 25,
    '学校': 18,
    '学院': 16,
    '采集人': 28,
    '采集日期': 14,
    '简介': 55,
    '主页': 45
}

for ws, data_list in [(ws1, with_email), (ws2, without_email)]:
    ws.views.sheetView[0].showGridLines = True
    
    # Headers
    ws.append(columns)
    ws.row_dimensions[1].height = 28
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 15)
        
    # Rows
    for row_idx, r in enumerate(data_list, 2):
        row_data = [r.get(c, "") for c in columns]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 36
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_name in ['序号', '姓名', '职称', '学校', '学院', '采集人', '采集日期']:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name in ['邮箱', '主页']:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_name == '简介':
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb.save(filepath)
print(f"Exported Excel to: {filepath}")

# Update global 教师列表.xlsx
global_file = "教师列表.xlsx"
if os.path.exists(global_file):
    try:
        global_wb = openpyxl.load_workbook(global_file)
        g_ws = global_wb.active
        existing_names = set()
        for row in g_ws.iter_rows(values_only=True, min_row=2):
            if row and row[1]:
                existing_names.add(str(row[1]).strip())
        
        current_max = g_ws.max_row
        appended = 0
        for r in records:
            if r['姓名'] not in existing_names:
                g_ws.append([current_max + appended, r['姓名'], r['职称'], r['邮箱'], r['学校'], r['学院'], r['采集人'], r['采集日期'], r['简介'], r['主页']])
                appended += 1
        global_wb.save(global_file)
        print(f"Appended {appended} new records to {global_file}")
    except Exception as e:
        print(f"Notice: Failed to update global {global_file}: {e}")

# Call Backend Sync
task_id = "1bbe14ee"
sync_url = f"http://127.0.0.1:8000/api/task/{task_id}/sync"
stats = {
    "total": len(records),
    "with_email": len(with_email),
    "without_email": len(without_email),
    "email_coverage": f"{(len(with_email)/len(records)*100):.1f}%" if records else "0.0%"
}
payload = {
    "records": records,
    "stats": stats,
    "download_url": f"/download/{filename}",
    "error": None,
    "progress": 100
}

req = urllib.request.Request(
    sync_url,
    data=json.dumps(payload, ensure_ascii=False).encode('utf-8'),
    headers={"Content-Type": "application/json; charset=utf-8"},
    method="POST"
)

try:
    with urllib.request.urlopen(req) as resp:
        res_data = json.loads(resp.read().decode('utf-8'))
        print("Backend sync response:", json.dumps(res_data, ensure_ascii=False, indent=2))
except Exception as e:
    print(f"Backend sync error: {e}")
