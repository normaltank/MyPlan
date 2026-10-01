"""Delivery receipts bind explicit affiliations to the leased college."""
import json

import openpyxl
import pytest

from server.agent_queue import AgentTaskQueue
from server.university_delivery import publish_delivery


def claim(queue):
    task = queue.submit_task('甲大学', '甲学院', batch_id='batch')
    leased = queue.claim_task('worker', task_id=task.task_id)
    return task, leased


def record(**fields):
    return {'序号': 1, '姓名': '合成人名', '职称': '教授', '邮箱': '',
            '学校': '甲大学', '学院': '甲学院', '采集人': 'test',
            '采集日期': '2026-10-01', '简介': '完整履历',
            '主页': 'https://example.edu/profile', **fields}


@pytest.mark.parametrize('fields', [
    {'学校': '乙大学'}, {'学院': '乙学院'},
    {'school': '乙大学'}, {'college': '乙学院'},
    {'学校': '甲大学', 'school': '乙大学'},
    {'学院': '甲学院', 'college': '乙学院'},
    {'学校': '', 'school': '乙大学'},
])
def test_foreign_affiliation_cannot_complete_or_persist_the_college(tmp_path, fields):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task, leased = claim(queue)
    with pytest.raises(ValueError, match='scope'):
        queue.complete_task(task.task_id, [record(**fields)],
                            {'total': 1, 'with_email': 0, 'no_email': 1},
                            worker_id='worker', lease_token=leased.lease_token)
    saved = queue.get_task(task.task_id)
    assert saved.status == 'running' and saved.records == []


def test_scope_less_legacy_profileless_receipt_stays_supported(tmp_path):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task, leased = claim(queue)
    legacy = {'name': '合成人名', 'email': ''}
    completed = queue.complete_task(task.task_id, [legacy],
                                    {'total': 1, 'with_email': 0, 'no_email': 1},
                                    worker_id='worker', lease_token=leased.lease_token)
    assert completed.status == 'completed' and completed.records == [legacy]


@pytest.mark.parametrize('bio_key', ['简介', 'bio', 'full_bio'])
def test_oversize_biography_cannot_complete_or_persist(tmp_path, bio_key):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task, leased = claim(queue)
    with pytest.raises(ValueError, match='excel_cell_too_long'):
        queue.sync_task_result(task.task_id, [{'name': '合成人名', bio_key: '甲' * 32768}],
                               {'total': 1, 'with_email': 0, 'no_email': 1}, '/download/unused.xlsx',
                               worker_id='worker', lease_token=leased.lease_token)
    saved = queue.get_task(task.task_id)
    assert saved.status == 'running' and saved.records == [] and not saved.download_url


def test_matching_english_affiliations_and_limit_length_are_preserved(tmp_path):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task, leased = claim(queue)
    legacy = {'name': '合成人名', 'school': '甲大学', 'college': '甲学院', 'full_bio': '甲' * 32767}
    completed = queue.complete_task(task.task_id, [legacy],
                                    {'total': 1, 'with_email': 0, 'no_email': 1},
                                    worker_id='worker', lease_token=leased.lease_token)
    assert completed.records == [legacy]


@pytest.mark.parametrize('invalid', ['\ud800', '\ufffe', '\uffff'])
def test_invalid_xml_biography_cannot_complete_or_persist(tmp_path, invalid):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task, leased = claim(queue)
    with pytest.raises(ValueError, match='excel_invalid_xml_character'):
        queue.sync_task_result(task.task_id, [{'name': '合成人名', 'bio': f'原文{invalid}履历'}],
                               {'total': 1, 'with_email': 0, 'no_email': 1}, '/download/unused.xlsx',
                               worker_id='worker', lease_token=leased.lease_token)
    saved = queue.get_task(task.task_id)
    assert saved.status == 'running' and saved.records == [] and not saved.download_url


@pytest.mark.parametrize('email', [None, '', '   '])
def test_university_manifest_matches_empty_email_sheet(tmp_path, email):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task, leased = claim(queue)
    queue.complete_task(task.task_id, [record(邮箱=email)],
                        {'total': 1, 'with_email': 0, 'no_email': 1},
                        worker_id='worker', lease_token=leased.lease_token)
    delivery = publish_delivery(queue, task.batch_id, task.school, tmp_path / 'outputs')
    with open(delivery['manifest_path'], encoding='utf-8') as receipt:
        manifest = json.load(receipt)
    assert (delivery['total'], delivery['with_email'], delivery['no_email']) == (1, 0, 1)
    assert (manifest['total'], manifest['with_email'], manifest['no_email']) == (1, 0, 1)
    book = openpyxl.load_workbook(delivery['excel_path'])
    assert book['教师列表'].max_row == 1 and book['无邮箱教师列表'].max_row == 2
    book.close()
