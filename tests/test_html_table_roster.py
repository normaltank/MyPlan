"""Mixed staff tables: synthetic reproductions of case-library DOM shapes.

The original ahstu_dk/hbnu_wd browser snapshots live in the cloud evidence
directory, not Git. These fixtures use the saved plans, not live-site counts.
"""

import json
from pathlib import Path
from urllib.parse import urljoin

import pytest
from openpyxl import load_workbook

from server.browser_batch_tool import run_plan
from server.harvest_workflow import CapturedRosterPage, reconcile_rosters


PLANS = Path(__file__).parents[1] / "docs/research/case-library/plans"


def case_page(case, table):
    plan = json.loads((PLANS / f"{case}.json").read_text(encoding="utf-8"))["plan"]
    branch = plan["branches"][0]
    wrapper = "tableContainer" if case == "ahstu_dk" else "wp_articlecontent"
    return plan, CapturedRosterPage(branch["url"], f'<div class="{wrapper}">{table}</div>', branch)


def test_remark_column_keeps_every_staff_role_as_row_evidence():
    plan, page = case_page("ahstu_dk", """
        <table><tbody>
        <tr class="firstRow"><td>序号</td><td>姓名</td><td>职称</td><td>备注</td></tr>
        <tr><td>1</td><td>张三</td><td>教授</td><td>专任教师</td></tr>
        <tr><td>2</td><td>李四</td><td></td><td>党政办公室</td></tr>
        <tr><td>3</td><td>王五</td><td>实验师</td><td>实验教辅</td></tr>
        </tbody></table>""")
    result = reconcile_rosters([page], plan)
    assert result.unresolved == []
    assert [person.name for person in result.candidates] == ["张三", "李四", "王五"]
    for person, remark in zip(result.candidates, ["专任教师", "党政办公室", "实验教辅"]):
        assert person.table_row["cells"]["备注"] == remark
        assert person.table_row["section"] is None


def test_sections_and_repeated_headers_are_not_people():
    header = "<tr><td>部门</td><td>姓名</td><td>职称</td><td>邮箱</td><td>备注</td></tr>"
    plan, page = case_page("hbnu_wd", "<table><tbody>"
        '<tr><td colspan="5">一、专任教师</td></tr>' + header +
        '<tr><td>物理系</td><td>张三</td><td>教授</td><td></td><td></td></tr>' + header +
        '<tr><td>物理系</td><td>王五</td><td>讲师</td><td></td><td></td></tr>' +
        '<tr><td colspan="5">四、行政工作人员</td></tr>' + header +
        '<tr><td>办公室</td><td>李四</td><td></td><td></td><td></td></tr>' +
        "</tbody></table>")
    result = reconcile_rosters([page], plan)
    assert result.unresolved == []
    assert [person.name for person in result.candidates] == ["张三", "王五", "李四"]
    assert [person.table_row["section"] for person in result.candidates] == [
        "一、专任教师", "一、专任教师", "四、行政工作人员"]
    assert result.observations[0]["visible_cards"] == 3


def test_rowspan_remark_and_split_name_stay_with_each_person():
    plan, page = case_page("ahstu_dk", """
        <table><tbody>
        <tr class="firstRow"><th>姓名</th><th>职称</th><th>备注</th></tr>
        <tr><td><span>张</span><span>三</span></td><td>讲师</td><td rowspan="2">行政工作人员</td></tr>
        <tr><td>李四</td><td>助教</td></tr>
        <tr><td>王五</td><td>教授</td><td>特聘教师</td></tr>
        </tbody></table>""")
    result = reconcile_rosters([page], plan)
    assert result.unresolved == []
    assert [person.name for person in result.candidates] == ["张三", "李四", "王五"]
    assert [person.table_row["cells"]["备注"] for person in result.candidates] == [
        "行政工作人员", "行政工作人员", "特聘教师"]


def test_sections_reset_at_table_boundary_and_do_not_change_titles():
    plan, page = case_page("ahstu_dk", """
        <table><tbody><tr><th colspan="3">实验技术人员</th></tr>
        <tr class="firstRow"><td>姓 名</td><td>职称</td><td>备注</td></tr>
        <tr><td>张三</td><td>讲师</td><td>实验室</td></tr>
        <tr class="firstRow"><td>姓 名</td><td>职称</td><td>备注</td></tr>
        <tr><td>李四</td><td>教授</td><td>实验室</td></tr></tbody></table>
        <table><tbody><tr class="firstRow"><td>姓名</td><td>职称</td><td>备注</td></tr>
        <tr><td>王五</td><td>副教授</td><td></td></tr></tbody></table>""")
    result = reconcile_rosters([page], plan)
    assert result.unresolved == []
    assert [(person.name, person.title, person.table_row["section"]) for person in result.candidates] == [
        ("张三", "讲师", "实验技术人员"), ("李四", "教授", "实验技术人员"), ("王五", "副教授", None)]
    assert [person.table_row["table_index"] for person in result.candidates] == [1, 1, 2]


def test_name_column_prevents_another_cell_or_link_from_becoming_a_person():
    plan, page = case_page("ahstu_dk", """
        <table><tbody>
        <tr class="firstRow"><td>联系人</td><td>姓名</td><td>备注</td></tr>
        <tr><td><a href="/office">王五</a></td><td><a href="/person">张三</a></td><td>办公室</td></tr>
        </tbody></table>""")
    result = reconcile_rosters([page], plan)
    assert result.unresolved == []
    assert [(person.name, person.url) for person in result.candidates] == [
        ("张三", "https://www.ahstu.edu.cn/person")]
    assert result.candidates[0].table_row["cells"]["联系人"] == "王五"


def test_unknown_name_in_an_administrative_section_still_needs_review():
    plan, page = case_page("ahstu_dk", """
        <table><tbody><tr><td colspan="2">行政工作人员</td></tr>
        <tr class="firstRow"><td>姓名</td><td>备注</td></tr>
        <tr><td>维诺尔科娃</td><td>办公室</td></tr></tbody></table>""")
    result = reconcile_rosters([page], plan)
    assert result.candidates == []
    assert [item["status"] for item in result.unresolved] == ["rare_surname_needs_review"]
    evidence = result.observations[0]["html_table_rows"][0]
    assert evidence["name"] == "维诺尔科娃" and evidence["section"] == "行政工作人员"


def test_partial_colspan_person_and_empty_spacer_are_not_section_boundaries():
    plan, page = case_page("ahstu_dk", """
        <table><tbody><tr><td colspan="4">行政工作人员</td></tr>
        <tr class="firstRow"><td>姓名</td><td>学位</td><td>职称</td><td>备注</td></tr>
        <tr><td colspan="2">张三</td><td rowspan="2">讲师</td><td rowspan="2">办公室</td></tr>
        <tr><td colspan="2">李四</td></tr>
        <tr><td colspan="4">&nbsp;</td></tr>
        <tr><td>王五</td><td>博士</td><td>教授</td><td></td></tr></tbody></table>""")
    result = reconcile_rosters([page], plan)
    assert result.unresolved == []
    assert [person.name for person in result.candidates] == ["张三", "李四", "王五"]
    assert [person.table_row["section"] for person in result.candidates] == ["行政工作人员"] * 3
    assert result.candidates[1].table_row["cells"]["备注"] == "办公室"
    assert result.observations[0]["blank_cards"] == 1


@pytest.mark.asyncio
@pytest.mark.browser
@pytest.mark.parametrize("case", ["ahstu_dk", "hbnu_wd"])
async def test_batch_delivery_keeps_row_evidence_and_standard_dual_sheets(tmp_path, case):
    from playwright.async_api import async_playwright

    table = """
        <table><tbody><tr class="firstRow"><td>姓名</td><td>职称</td><td>备注</td></tr>
        <tr><td>张三</td><td>教授</td><td>专任教师</td></tr>
        <tr><td>李四</td><td></td><td>党政办公室</td></tr>
        <tr><td>王五</td><td>实验师</td><td>实验教辅</td></tr></tbody></table>"""
    if case == "hbnu_wd":
        header = "<tr><td>部门</td><td>姓名</td><td>职称</td><td>邮箱</td><td>备注</td></tr>"
        table = ('<table><tbody><tr><td colspan="5">一、专任教师</td></tr>' + header +
                 '<tr><td>物理系</td><td>张三</td><td>教授</td><td></td><td>专任教师</td></tr>' +
                 '<tr><td colspan="5">四、行政工作人员</td></tr>' + header +
                 '<tr><td>办公室</td><td>李四</td><td></td><td></td><td>党政办公室</td></tr>' + header +
                 '<tr><td colspan="5">实验技术人员</td></tr>' + header +
                 '<tr><td>实验室</td><td>王五</td><td>实验师</td><td></td><td>实验教辅</td></tr>' +
                 '</tbody></table>')
    plan, page = case_page(case, table)
    # Both a supplementary table and a unique personal page name 张三.
    # The existing identity merge must retain the table's role evidence.
    linked_url = urljoin(page.url, "/linked")
    profile_url = urljoin(page.url, "/person")
    plan["branches"].append({"url": linked_url, "card_selector": "li", "single_page": True})
    pages = {
        page.url: page.html,
        linked_url: '<ul><li><a href="/person">张三</a></li></ul>',
        profile_url: '<article><h1>张三</h1><p>教授，完整科研履历原文。</p>'
                     '<p>邮箱：zhangsan@example.edu.cn</p></article>',
    }

    async def serve(route):
        if route.request.url in pages:
            await route.fulfill(status=200, content_type="text/html; charset=utf-8", body=pages[route.request.url])
        else:
            await route.abort()

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True)
        context = await browser.new_context()
        await context.route("**/*", serve)
        report = await run_plan(plan, context, tmp_path)
        await browser.close()
    assert report["delivery_ready"] is True and report["unresolved"] == []
    assert report["total"] == 3 and report["with_email"] == 1 and report["no_email"] == 2
    rows = json.loads(Path(report["records_path"]).read_text(encoding="utf-8"))
    coverage = json.loads(Path(report["coverage_path"]).read_text(encoding="utf-8"))
    for source in coverage["field_sources"]:
        row = rows[source["record_index"] - 1]
        evidence, = source["html_table_rows"]
        assert evidence["name"] == row["姓名"]
        assert evidence["roster_url"] == page.url
        # Chromium adds html/head/body around the fixture before capture.
        assert page.html in Path(evidence["roster_snapshot"]).read_text(encoding="utf-8")
        if case == "hbnu_wd":
            assert evidence["section"] == {"张三": "一、专任教师", "李四": "四、行政工作人员", "王五": "实验技术人员"}[row["姓名"]]
        assert evidence["cells"]["备注"] == {"张三": "专任教师", "李四": "党政办公室", "王五": "实验教辅"}[row["姓名"]]
    assert next(row for row in rows if row["姓名"] == "张三")["简介"] == (
        "张三\n教授，完整科研履历原文。\n邮箱：zhangsan@example.edu.cn")
    workbook = load_workbook(report["excel_path"])
    assert workbook.sheetnames == ["教师列表", "无邮箱教师列表"]
    assert sum(sheet.max_row - 1 for sheet in workbook) == 3
    for sheet in workbook:
        assert [cell.value for cell in sheet[1]] == [
            "序号", "姓名", "职称", "邮箱", "学校", "学院", "采集人", "采集日期", "简介", "主页"]
