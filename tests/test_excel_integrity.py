"""The shared export interface must not silently lose harvested text."""
import openpyxl
import pytest

from server.excel_export import export_dual_sheet_excel


def record(bio):
    return {'序号': 1, '姓名': '张三', '职称': '教授', '邮箱': '', '学校': '测试大学',
            '学院': '测试学院', '采集人': 'fixture', '采集日期': '2026-10-01',
            '简介': bio, '主页': 'https://example.edu/person'}


def test_overlong_biography_is_preserved_in_the_workbook(tmp_path):
    destination = tmp_path / 'university.xlsx'
    export_dual_sheet_excel([record('完整旧简介')], str(destination))
    original = '文' * 32768
    counts = export_dual_sheet_excel([record(original)], str(destination))
    workbook = openpyxl.load_workbook(destination)
    parts = [row[9] for row in list(workbook['完整简介'].values)[1:]]
    assert ''.join(parts) == original
    assert counts['total'] == 1 and counts['full_biographies'] == 1
    assert workbook['无邮箱教师列表']['I2'].hyperlink.location == "'完整简介'!A2"


def test_failed_workbook_serialization_keeps_the_previous_delivery(tmp_path):
    destination = tmp_path / 'university.xlsx'
    export_dual_sheet_excel([record('完整旧简介')], str(destination))
    before = destination.read_bytes()
    with pytest.raises(ValueError):
        export_dual_sheet_excel([record('原文包含 XML 不可表示字符\uffff')], str(destination))
    assert destination.read_bytes() == before
    assert sorted(path.name for path in tmp_path.iterdir()) == ['university.xlsx']


def test_failed_atomic_publish_keeps_previous_file_and_cleans_temporary(tmp_path, monkeypatch):
    from server import excel_export
    destination = tmp_path / 'university.xlsx'
    export_dual_sheet_excel([record('完整旧简介')], str(destination))
    before = destination.read_bytes()

    def unavailable(*_args):
        raise PermissionError('destination is locked')

    monkeypatch.setattr(excel_export.os, 'replace', unavailable)
    with pytest.raises(PermissionError, match='locked'):
        export_dual_sheet_excel([record('完整新简介')], str(destination))
    assert destination.read_bytes() == before
    assert sorted(path.name for path in tmp_path.iterdir()) == ['university.xlsx']


def test_cell_limit_preserves_all_characters_and_dual_sheet_partition(tmp_path):
    destination = tmp_path / 'university.xlsx'
    original = '文' * 32767
    counts = export_dual_sheet_excel([record(original), {**record('完整履历'), '姓名': '李四',
                                                      '邮箱': 'li@example.edu'}], str(destination))
    workbook = openpyxl.load_workbook(destination)
    assert workbook['无邮箱教师列表']['I2'].value == original
    assert counts == {'total': 2, 'with_email': 1, 'no_email': 1}
    assert workbook['教师列表']['B2'].value == '李四'


def test_public_export_preserves_legacy_fields_in_standard_columns(tmp_path):
    destination = tmp_path / 'university.xlsx'
    biography = '  完整原文\n\n包括结尾空白。  \n'
    raw = {'name': '张三', 'title': '教授', 'email': 'zhang@example.edu',
           'school': '测试大学', 'college': '测试学院', 'full_bio': biography,
           'profile_url': 'https://example.edu/person'}
    counts = export_dual_sheet_excel([raw], str(destination))
    workbook = openpyxl.load_workbook(destination)
    row = list(workbook['教师列表'].values)[1]
    assert row[1:6] == ('张三', '教授', 'zhang@example.edu', '测试大学', '测试学院')
    assert row[8] == biography
    assert counts == {'total': 1, 'with_email': 1, 'no_email': 0}
    assert raw['full_bio'] == biography and '简介' not in raw


def test_legacy_affiliation_alias_is_checked_before_excel_can_truncate_it(tmp_path):
    destination = tmp_path / 'university.xlsx'
    with pytest.raises(ValueError, match='excel_cell_too_long'):
        export_dual_sheet_excel([{**record('完整简介'), '学校': '', 'school': 'U' * 32768}],
                               str(destination))
    assert not destination.exists()


def test_empty_workbook_save_failure_cannot_publish_a_partial_file(tmp_path, monkeypatch):
    from pathlib import Path
    from server import excel_export
    destination = tmp_path / 'university.xlsx'
    export_dual_sheet_excel([record('旧交付完整原文')], str(destination))
    before = destination.read_bytes()
    original_load = excel_export.openpyxl.load_workbook

    def load_with_failing_save(*args, **kwargs):
        workbook = original_load(*args, **kwargs)

        def failed_save(path):
            Path(path).write_bytes(b'partial ZIP')
            raise OSError('disk write failed')

        workbook.save = failed_save
        return workbook

    monkeypatch.setattr(excel_export.openpyxl, 'load_workbook', load_with_failing_save)
    with pytest.raises(OSError, match='disk write failed'):
        export_dual_sheet_excel([], str(destination))
    assert destination.read_bytes() == before
    assert sorted(path.name for path in tmp_path.iterdir()) == ['university.xlsx']


@pytest.mark.asyncio
async def test_harvest_keeps_full_overlong_bio_and_exports_every_character(tmp_path):
    import json
    from pathlib import Path
    from server import browser_batch_tool as batch
    from test_browser_batch_tool import FakeContext, roster

    base = 'https://example.edu'
    original = '完整研究经历。' * 6000
    html = f'<html><body><article><h1>张富有</h1><p>{original}</p></article></body></html>'
    context = FakeContext({base + '/roster': roster([(1, '张富有')]), base + '/p/1.htm': html})
    plan = {'school': '测试大学', 'college': '测试学院', 'branches': [
        {'url': base + '/roster', 'card_selector': 'main li', 'single_page': True}]}

    report = await batch.run_plan(plan, context, tmp_path)

    assert report['delivery_ready'] is True
    rows = json.loads(Path(report['records_path']).read_text(encoding='utf-8'))
    assert original in rows[0]['简介']
    assert len(rows[0]['简介']) > 32767
    issues = [item for item in report['unresolved'] if item['status'] == 'excel_cell_too_long']
    assert not issues
    workbook = openpyxl.load_workbook(report['excel_path'])
    assert ''.join(row[9] for row in list(workbook['完整简介'].values)[1:]) == rows[0]['简介']

# This component fixture stipulates verified source purpose; purpose integration is unmocked elsewhere.
pytestmark = pytest.mark.verified_source_boundary
