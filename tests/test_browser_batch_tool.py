"""Browser-backed college batches must expose gaps instead of silently exporting."""

import asyncio
import json
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import openpyxl
import pytest

import server.browser_batch_tool as batch_tool
from server.browser_batch_tool import _capture_branch, _roster_signature, run_plan
from server.faculty_extractor import FacultyCandidate, FacultyHarvestPipeline
from server.harvest_workflow import CapturedRosterPage, reconcile_rosters


def roster(names):
    cards = "".join(
        f'<li><a href="/p/{i}.htm">{name}</a><span>教授</span></li>'
        for i, name in names
    )
    return f"<html><body><main><ul>{cards}</ul></main></body></html>"


def profile(name, email):
    return (
        f"<html><body><article><h1>{name}</h1>"
        f"<p>邮箱：{email}</p><p>完整履历：{name}的学习经历、科研项目、论文甲。</p>"
        "</article></body></html>"
    )


class FakeLocator:
    def __init__(self, page, selector):
        self.page = page
        self.selector = selector

    async def count(self):
        if self.selector != ".next":
            return 0
        return int(self.page.index < len(self.page.context.pages[self.page.url]) - 1)

    async def is_enabled(self):
        return True

    async def click(self):
        if self.page.url not in self.page.context.stuck:
            self.page.index += 1


class FakePage:
    def __init__(self, context):
        self.context = context
        self.url = ""
        self.index = 0

    async def goto(self, url, **_kwargs):
        self.url = url
        self.index = 0
        self.context.visits.append(url)
        if url not in self.context.pages:
            raise RuntimeError("browser navigation failed")

        class Response:
            status = 200

        return Response()

    async def content(self):
        return self.context.pages[self.url][self.index]

    def locator(self, selector):
        return FakeLocator(self, selector)

    async def close(self):
        pass


class FakeContext:
    def __init__(self, pages, stuck=(), assets=None):
        self.pages = {url: [html] if isinstance(html, str) else html for url, html in pages.items()}
        self.stuck = set(stuck)
        self.visits = []
        self.assets = assets or {}

        class Request:
            async def get(_, url, **_kwargs):
                if url not in self.assets:
                    raise RuntimeError("asset unavailable")

                class Response:
                    ok = True

                    async def body(inner):
                        return self.assets[url]

                return Response()

        self.request = Request()

    async def new_page(self):
        return FakePage(self)


@pytest.mark.asyncio
async def test_single_page_assertion_is_rejected_when_second_roster_page_exists():
    base = "https://faculty.example.edu/roster"
    first = roster([(1, "张富有")]).replace(
        "</main>", '<nav class="pagination"><a href="/roster?page=2">2</a></nav></main>')
    second = roster([(2, "刘圣洁")])
    browser = FakeContext({base: first, "https://faculty.example.edu/roster?page=2": second})

    pages, errors, evidence = await _capture_branch(
        browser, {"url": base, "single_page": True, "card_selector": "main li"})

    assert len(pages) == 1
    assert any(item["status"] == "single_page_has_next" for item in errors)
    assert evidence["end"] != "agent_asserted_single_page"


def test_missing_card_selector_blocks_roster_reconciliation():
    url = "https://faculty.example.edu/roster"
    captured = [CapturedRosterPage(url, roster([(1, "张富有")]), {"url": url})]

    result = reconcile_rosters(captured, {"school": "样本大学", "college": "样本学院"})

    assert any(item["status"] == "card_selector_missing" for item in result.unresolved)


@pytest.mark.asyncio
async def test_single_page_assertion_rejects_ajax_next_button():
    url = "https://faculty.example.edu/roster"
    html = roster([(1, "张富有")]).replace(
        "</main>", '<button class="next" type="button">下一页</button></main>')
    browser = FakeContext({url: [html, roster([(2, "刘圣洁")])]})

    _, errors, _ = await _capture_branch(
        browser, {"url": url, "single_page": True, "card_selector": "main li"})

    assert any(item["status"] == "single_page_pagination_unverified" for item in errors)


@pytest.mark.asyncio
async def test_same_url_ajax_anchor_cannot_pass_single_page_audit():
    url = "https://faculty.example.edu/roster"
    html = roster([(1, "张富有")]).replace(
        "</main>", f'<a class="next" href="{url}" onclick="event.preventDefault()">下一页</a></main>')
    _, errors, _ = await _capture_branch(
        FakeContext({url: html}), {"url": url, "single_page": True, "card_selector": "main li"})
    assert any(item["status"] == "single_page_pagination_unverified" for item in errors)


@pytest.mark.asyncio
async def test_body_page_class_does_not_make_every_profile_link_a_pager():
    url = "https://faculty.example.edu/roster"
    html = roster([(1, "张富有")]).replace("<body>", '<body class="page">')

    _, errors, _ = await _capture_branch(
        FakeContext({url: html}),
        {"url": url, "single_page": True, "card_selector": "main li"})

    assert errors == []


@pytest.mark.asyncio
async def test_single_page_probe_does_not_follow_external_pager_link():
    url = "https://faculty.example.edu/roster"
    html = roster([(1, "张富有")]).replace(
        "</main>", '<nav class="pagination"><a href="https://other.example/page2">2</a></nav></main>')
    browser = FakeContext({url: html})

    _, errors, _ = await _capture_branch(
        browser, {"url": url, "single_page": True, "card_selector": "main li"})

    assert any(item["status"] == "single_page_pagination_unverified" for item in errors)
    assert browser.visits == [url]


@pytest.mark.asyncio
async def test_native_browser_recovery_does_not_run_sync_playwright_on_event_loop(monkeypatch):
    url = "https://faculty.example.edu/blocked"

    def recovered_html(_url):
        with pytest.raises(RuntimeError, match="no running event loop"):
            asyncio.get_running_loop()
        return roster([(1, "张富有")]) + " " * 1100

    monkeypatch.setattr(batch_tool, "_fetch_page_via_native_chrome", recovered_html)
    pages, errors, _ = await _capture_branch(
        FakeContext({}), {"url": url, "single_page": True, "card_selector": "main li"})

    assert len(pages) == 1
    assert errors == []


@pytest.mark.asyncio
async def test_browser_batch_captures_all_pages_profiles_and_full_bio(tmp_path):
    base = "https://faculty.example.edu"
    first = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    second = [(4, "赵文成"), (5, "王英华"), (6, "周立新")]
    pages = {f"{base}/roster": [roster(first), roster(second)]}
    pages.update({f"{base}/p/{i}.htm": profile(name, f"{i:010d}@example.edu")
                  for i, name in first + second})
    browser = FakeContext(pages)
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "next_selector": ".next", "max_pages": 3,
         "card_selector": "main li"}
    ]}

    result = await run_plan(plan, browser, tmp_path)

    assert result["delivery_ready"] is True
    assert result["total"] == 6
    assert result["with_email"] == 6
    assert result["directory_pages"] == 2
    assert result["profile_pages"] == 6
    assert result["profile_attempts"] == 6
    assert result["branches"][0]["next_clicks"] == 1
    assert result["elapsed_seconds"] >= 0
    assert len(browser.visits) == 7
    assert result["unresolved"] == []
    records = json.loads(Path(result["records_path"]).read_text(encoding="utf-8"))
    assert len(records) == 6
    assert all("学习经历、科研项目、论文甲" in row["简介"] for row in records)
    assert all(set(row) == {"序号", "姓名", "职称", "邮箱", "学校", "学院", "采集人", "采集日期", "简介", "主页"}
               for row in records)
    assert len(result["snapshots"]) == 8
    assert all(Path(item["path"]).is_file() for item in result["snapshots"])
    workbook = openpyxl.load_workbook(result["excel_path"], read_only=True)
    assert workbook.sheetnames == ["教师列表", "无邮箱教师列表"]
    assert workbook["教师列表"].max_row == 7
    assert workbook["无邮箱教师列表"].max_row == 1
    assert len(list(workbook["教师列表"].values)[0]) == 10


@pytest.mark.asyncio
async def test_image_biography_is_added_from_original_asset_with_ocr_evidence(tmp_path, monkeypatch):
    import server.image_biography as media

    base = "https://faculty.example.edu"
    url = f"{base}/p/1.htm"
    image_url = f"{base}/cv.png"
    browser = FakeContext({f"{base}/roster": roster([(1, "张富有")]),
                           url: '<html><body><article><h1>张富有</h1><img src="/cv.png"></article></body></html>'},
                          assets={image_url: b"original image"})
    monkeypatch.setattr(media, "scan_media", lambda *_args, **_kwargs: {
        "page_count": 1, "pages_ocrd": 1, "text": "张富有 教授\n完整学习经历与任职信息\n研究项目与学术成果详述\n邮箱 zhang@example.edu",
        "pages": [{"page": 1, "text": "张富有 教授\n完整学习经历与任职信息\n研究项目与学术成果详述\n邮箱 zhang@example.edu",
                   "lines": [{"text": line, "confidence": 0.99} for line in
                             ("张富有 教授", "完整学习经历与任职信息", "研究项目与学术成果详述", "邮箱 zhang@example.edu")],
                   "source": "local_ocr"}], "image_size": [1200, 1800]})
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "single_page": True, "card_selector": "main li"}]}

    result = await run_plan(plan, browser, tmp_path)

    assert result["delivery_ready"] is True
    assert len(result["image_biography_evidence"]) == 1
    record = json.loads(Path(result["records_path"]).read_text(encoding="utf-8"))[0]
    assert "完整学习经历" in record["简介"]
    assert record["邮箱"] == "zhang@example.edu"
    assert result["with_email"] == 1
    assert result["coverage"]["field_sources"][0]["extraction"] == "deterministic_dom_and_local_ocr"


@pytest.mark.asyncio
async def test_missing_image_biography_asset_blocks_delivery(tmp_path):
    base = "https://faculty.example.edu"
    url = f"{base}/p/1.htm"
    browser = FakeContext({f"{base}/roster": roster([(1, "张富有")]),
                           url: '<html><body><article><h1>张富有</h1><img src="/cv.png"></article></body></html>'})
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "single_page": True, "card_selector": "main li"}]}

    result = await run_plan(plan, browser, tmp_path)

    assert result["delivery_ready"] is False
    assert result["excel_path"] is None
    assert any(item["status"] == "image_bio_asset_failed" for item in result["unresolved"])


@pytest.mark.asyncio
async def test_image_only_profile_without_usable_ocr_blocks_delivery(tmp_path, monkeypatch):
    import server.image_biography as media

    base = "https://faculty.example.edu"
    url = f"{base}/p/1.htm"
    browser = FakeContext({f"{base}/roster": roster([(1, "张富有")]),
                           url: '<html><body><article><img src="/cv.png"></article></body></html>'},
                          assets={f"{base}/cv.png": b"original image"})
    monkeypatch.setattr(media, "scan_media", lambda *_args, **_kwargs: {
        "page_count": 1, "pages_ocrd": 1, "text": "", "pages": [
            {"page": 1, "text": "", "lines": [], "source": "local_ocr"}],
        "image_size": [400, 600]})
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "single_page": True, "card_selector": "main li",
         "profile_content_selector": "article"}]}

    result = await run_plan(plan, browser, tmp_path)

    assert result["delivery_ready"] is False
    assert result["excel_path"] is None
    assert any(item["status"] == "image_bio_unreadable_profile" for item in result["unresolved"])


@pytest.mark.asyncio
async def test_roster_inside_misused_header_is_still_captured(tmp_path):
    base = "https://faculty.example.edu"
    names = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    html = "<html><body><header><nav>网站导航</nav><main><ul class='faculty'>" + "".join(
        f'<li><a href="/p/{number}.htm">{person}</a><span>教授</span></li>'
        for number, person in names
    ) + "</ul></main></header></body></html>"
    pages = {f"{base}/roster": html}
    pages.update({f"{base}/p/{number}.htm": profile(person, f"{number:010d}@example.edu")
                  for number, person in names})
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "single_page": True, "card_selector": "main li"}
    ]}

    result = await run_plan(plan, FakeContext(pages), tmp_path)

    assert result["delivery_ready"] is True
    assert result["total"] == 3


@pytest.mark.asyncio
async def test_hospital_snapshot_accounts_for_cards_but_still_requires_eligibility(tmp_path):
    """Structural recall accounts for cards; hospital employment still needs proof."""
    snapshot = next((Path(__file__).parents[1] / "server" / "outputs" /
                     "甘肃中医药大学_第二临床医学院_snapshots").glob("directory_001*.html"))
    url = "http://www.lzsdyyy.cn/html/doctors-list-lingjunrenwu.html"
    plan = {"school": "甘肃中医药大学", "college": "第二临床医学院", "branches": [
        {"url": url, "single_page": True, "card_selector": "ul.expert-list > li"}
    ]}

    result = await run_plan(plan, FakeContext({url: snapshot.read_text(encoding="utf-8")}), tmp_path)

    assert result["delivery_ready"] is False
    observation = result["roster_observations"][0]
    assert observation["visible_cards"] == 120
    assert observation["extracted_candidates"] == 120
    assert observation["unaccounted_cards"] == []
    assert not any(item["status"] == "roster_card_mismatch" for item in result["unresolved"])
    assert result["unresolved"]


@pytest.mark.asyncio
async def test_navigation_page_discovers_bounded_roster_branches(tmp_path):
    base = "https://faculty.example.edu"
    root = f"{base}/departments"
    pages = {root: (
        '<html><body><section class="departments">'
        '<a href="/department/a">系所 A</a><a href="/department/b">系所 B</a>'
        '</section></body></html>'
    )}
    first = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    second = [(4, "赵文成"), (5, "王英华"), (6, "周立新")]
    pages[f"{base}/department/a"] = roster(first)
    pages[f"{base}/department/b"] = roster(second)
    pages.update({f"{base}/p/{number}.htm": profile(person, f"{number:010d}@example.edu")
                  for number, person in first + second})
    plan = {"school": "样本大学", "college": "样本学院", "branches": [{
        "url": root, "kind": "navigation", "child_link_selector": ".departments a",
        "max_children": 3, "child_branch": {"single_page": True, "card_selector": "main li"},
    }]}

    result = await run_plan(plan, FakeContext(pages), tmp_path)

    assert result["delivery_ready"] is True
    assert result["navigation_pages"] == 1
    assert result["directory_pages"] == 2
    assert result["total"] == 6
    assert result["branches"][0]["discovered_children"] == 2
    assert all(item["kind"] != "directory" for item in result["snapshots"][:1])


@pytest.mark.asyncio
async def test_hospital_doctor_roster_requires_individual_university_appointment_evidence(tmp_path):
    base = "https://hospital.example.cn"
    people = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    pages = {f"{base}/doctors": roster(people)}
    pages.update({f"{base}/p/{number}.htm": profile(person, "") for number, person in people})
    plan = {"school": "样本大学", "college": "样本学院",
            "source_kind": "hospital_doctor_roster", "branches": [
                {"url": f"{base}/doctors", "single_page": True, "card_selector": "main li"}
            ]}

    result = await run_plan(plan, FakeContext(pages), tmp_path)

    assert result["candidate_count"] == 3
    assert result["total"] == 0
    assert result["excluded_count"] == 3
    assert result["delivery_ready"] is False
    assert result["excel_path"] is None


@pytest.mark.asyncio
async def test_hospital_current_university_appointment_admits_only_proven_person(tmp_path):
    base = "https://hospital.example.cn"
    people = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    pages = {f"{base}/doctors": roster(people)}
    pages.update({f"{base}/p/{number}.htm": profile(person, "") for number, person in people})
    quote = "张富有现任样本大学样本学院教授"
    pages[f"{base}/p/1.htm"] = f"<html><body><article><h1>张富有</h1><p>{quote}</p></article></body></html>"
    plan = {"school": "样本大学", "college": "样本学院",
            "source_kind": "hospital_doctor_roster", "branches": [
                {"url": f"{base}/doctors", "single_page": True, "card_selector": "main li"}
            ], "verified_appointments": [
                {"name": "张富有", "profile_url": f"{base}/p/1.htm",
                 "quote": quote, "academic_title": "教授"}
            ]}

    result = await run_plan(plan, FakeContext(pages), tmp_path)

    assert result["delivery_ready"] is True
    assert result["candidate_count"] == 3
    assert result["excluded_count"] == 2
    assert result["total"] == 1
    record = json.loads(Path(result["records_path"]).read_text(encoding="utf-8"))[0]
    assert record["姓名"] == "张富有"
    assert record["职称"] == "教授"
    assert result["eligibility"][0]["evidence_kind"] == "hospital_current_university_appointment"


@pytest.mark.asyncio
async def test_unchanged_pagination_is_unresolved_and_blocks_export(tmp_path):
    url = "https://faculty.example.edu/roster"
    browser = FakeContext({url: [roster([(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]),
                                 roster([(4, "赵文成"), (5, "王英华"), (6, "周立新")])]}, stuck=[url])
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
                        {"url": url, "next_selector": ".next", "max_pages": 3,
                         "card_selector": "main li"}
    ]}

    result = await run_plan(plan, browser, tmp_path)

    assert result["delivery_ready"] is False
    assert any(item["status"] == "pagination_unchanged" for item in result["unresolved"])
    assert result["excel_path"] is None


@pytest.mark.asyncio
async def test_pagination_that_only_changes_page_chrome_is_unresolved(tmp_path):
    url = "https://faculty.example.edu/roster"
    names = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    first = roster(names)
    browser = FakeContext({url: [first + "<footer>clock 1</footer>",
                                 first + "<footer>clock 2</footer>"]})
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "next_selector": ".next", "max_pages": 3}
    ]}

    result = await run_plan(plan, browser, tmp_path)

    assert any(item["status"] == "pagination_unchanged" for item in result["unresolved"])
    assert result["directory_pages"] == 1


@pytest.mark.asyncio
async def test_profile_navigation_failure_is_unresolved_and_blocks_export(tmp_path):
    base = "https://faculty.example.edu"
    names = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    pages = {f"{base}/roster": roster(names)}
    pages.update({f"{base}/p/{i}.htm": profile(name, f"{i:010d}@example.edu")
                  for i, name in names[1:]})
    browser = FakeContext(pages)
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "single_page": True}
    ]}

    result = await run_plan(plan, browser, tmp_path)

    assert result["delivery_ready"] is False
    assert any(item["status"].startswith("fetch_failed") for item in result["unresolved"])
    assert result["excel_path"] is None


@pytest.mark.asyncio
async def test_empty_roster_needs_agent_review_instead_of_empty_delivery(tmp_path):
    url = "https://faculty.example.edu/roster"
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "single_page": True}
    ]}

    result = await run_plan(plan, FakeContext({url: "<html><body><main></main></body></html>"}), tmp_path)

    assert result["delivery_ready"] is False
    assert any(item["status"] == "no_candidates_on_roster_page" for item in result["unresolved"])
    assert result["excel_path"] is None


@pytest.mark.asyncio
async def test_rendered_profile_without_scholar_identity_requires_review(tmp_path):
    base = "https://faculty.example.edu"
    names = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    pages = {f"{base}/roster": roster(names)}
    pages.update({f"{base}/p/{i}.htm": profile(name, f"{i:010d}@example.edu")
                  for i, name in names[1:]})
    pages[f"{base}/p/1.htm"] = "<html><body><div id='app'>Loading...</div></body></html>"
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "single_page": True}
    ]}

    result = await run_plan(plan, FakeContext(pages), tmp_path)

    assert result["delivery_ready"] is False
    assert any(item["status"] == "profile_identity_not_found" for item in result["unresolved"])
    assert result["excel_path"] is None


@pytest.mark.asyncio
async def test_plan_must_assert_single_page_or_supply_next_control(tmp_path):
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": "https://faculty.example.edu/roster"}
    ]}
    with pytest.raises(ValueError, match="single_page|next_selector"):
        await run_plan(plan, FakeContext({}), tmp_path)


@pytest.mark.asyncio
@pytest.mark.parametrize("force_conflict", [False, True])
async def test_complete_shared_prose_and_genuine_extractor_conflicts(tmp_path, monkeypatch, force_conflict):
    base = "https://faculty.example.edu"
    names = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    pages = {f"{base}/roster": roster(names)}
    # Repeated body copy can be mistaken for a site-wide template by cluster extraction.
    pages.update({f"{base}/p/{i}.htm": (
        f'<html><body><div class="article"><h1>{name}</h1>'
        f'<p>邮箱：{i:010d}@example.edu</p><p>相同的完整履历正文。</p></div></body></html>'
    ) for i, name in names})
    browser = FakeContext(pages)
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": f"{base}/roster", "single_page": True, "card_selector": "main li"}
    ]}

    if force_conflict:
        original = FacultyHarvestPipeline.harvest_candidates
        def shortened(*args, **kwargs):
            result = original(*args, **kwargs)
            for row in result["sheet1_with_email"] + result["sheet2_no_email"]:
                row["简介"] = row["姓名"]
            return result
        monkeypatch.setattr(FacultyHarvestPipeline, "harvest_candidates", shortened)

    result = await run_plan(plan, browser, tmp_path)

    if force_conflict:
        assert result["delivery_ready"] is False
        assert result["excel_path"] is None
        assert any(item["status"] == "profile_body_disagreement" for item in result["unresolved"])
    else:
        assert result["delivery_ready"] is True
        assert result["total"] == result["with_email"] == 3
        records = json.loads(Path(result["records_path"]).read_text(encoding="utf-8"))
        assert all("相同的完整履历正文。" in row["简介"] for row in records)


@pytest.mark.asyncio
async def test_agent_verified_poster_names_can_be_added_with_image_evidence(tmp_path, monkeypatch):
    import server.image_biography as media
    url = "https://faculty.example.edu/poster"
    screenshot = tmp_path / "poster.png"
    screenshot.write_bytes(b"reviewed image")
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "single_page": True}
    ], "poster_reviews": [{"source_url": url, "reviewed_full_image": True, "expected_count": 2}],
    "verified_candidates": [
        {"name": "张富有", "title": "教授", "source_url": url, "evidence_path": str(screenshot)},
        {"name": "刘圣洁", "title": "副教授", "source_url": url, "evidence_path": str(screenshot)},
    ]}

    monkeypatch.setattr(media, "scan_media", lambda *_args, **_kwargs: {
        "page_count": 1, "pages_ocrd": 1, "text": "张富有 刘圣洁",
        "pages": [{"page": 1, "text": "张富有 刘圣洁", "source": "local_ocr",
                   "lines": [{"text": "张富有 刘圣洁", "confidence": 0.99}]}],
        "image_size": [400, 600]})
    result = await run_plan(plan, FakeContext(
        {url: "<html><body><img src='poster.png'></body></html>"},
        assets={"https://faculty.example.edu/poster.png": b"original poster"}), tmp_path)

    assert result["delivery_ready"] is True
    assert result["total"] == 2
    assert result["no_email"] == 2
    assert result["verified_candidate_count"] == 2
    assert result["unresolved"] == []


@pytest.mark.asyncio
async def test_verified_roster_images_add_separate_biographies(tmp_path, monkeypatch):
    import server.image_biography as media

    url = "https://faculty.example.edu/poster"
    image_a = "https://faculty.example.edu/a.png"
    image_b = "https://faculty.example.edu/b.png"
    screenshot = tmp_path / "reviewed.png"
    screenshot.write_bytes(b"reviewed")
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "single_page": True}],
        "poster_reviews": [{"source_url": url, "reviewed_full_image": True, "expected_count": 2}],
        "verified_candidates": [
            {"name": name, "source_url": url, "evidence_path": str(screenshot)}
            for name in ("张富有", "刘圣洁")]}
    html = '<html><body><article><img src="/a.png"><img src="/b.png"></article></body></html>'
    browser = FakeContext({url: html}, assets={image_a: b"a", image_b: b"b"})

    def scan(data, _kind, *, ocr_engine=None):
        name = "张富有" if data == b"a" else "刘圣洁"
        lines = (f"{name} 教授", "完整学习经历与任职信息", "研究项目与学术成果详述")
        text = "\n".join(lines)
        return {"page_count": 1, "pages_ocrd": 1, "text": text,
                "pages": [{"page": 1, "text": text, "source": "local_ocr",
                           "lines": [{"text": line, "confidence": 0.99} for line in lines]}],
                "image_size": [1200, 1800]}

    monkeypatch.setattr(media, "scan_media", scan)
    result = await run_plan(plan, browser, tmp_path)

    assert result["delivery_ready"] is True
    records = json.loads(Path(result["records_path"]).read_text(encoding="utf-8"))
    assert len(records) == 2
    assert all(row["姓名"] in row["简介"] for row in records)
    assert all(other not in row["简介"] for row in records
               for other in ("张富有", "刘圣洁") if other != row["姓名"])
    for row, source in zip(records, result["coverage"]["field_sources"]):
        assert len(source["image_biography_evidence"]) == 1
        assert source["image_biography_evidence"][0]["assigned_name"] == row["姓名"]


@pytest.mark.asyncio
async def test_human_selected_image_email_clears_only_matching_conflict(tmp_path, monkeypatch):
    from io import BytesIO
    from PIL import Image
    import server.image_biography as media
    from server.image_bio_review import collect_email_conflict_items, save_email_decision

    url = "https://faculty.example.edu/mentor"
    image_url = "https://faculty.example.edu/cv.png"
    original = BytesIO()
    Image.new("RGB", (500, 700), "white").save(original, format="PNG")
    browser = FakeContext({url: '<article><p>张富有</p><img src="/cv.png"></article>'},
                          assets={image_url: original.getvalue()})
    lines = ["张富有 教授", "完整学习经历与任职信息", "研究项目与学术成果详述",
             "邮箱 123456@qq.com", "邮箱 234567@163.com"]

    def scan(_data, _kind, *, ocr_engine=None):
        rows = [{"text": line, "confidence": .99,
                 "box": [[10, 20+i*30], [250, 20+i*30],
                         [250, 40+i*30], [10, 40+i*30]]}
                for i, line in enumerate(lines)]
        text = "\n".join(lines)
        return {"page_count": 1, "pages_ocrd": 1, "text": text,
                "pages": [{"page": 1, "text": text, "source": "local_ocr", "lines": rows}],
                "image_size": [500, 700]}

    monkeypatch.setattr(media, "scan_media", scan)
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "single_page": True, "card_selector": "article img"}]}
    first = await run_plan(plan, browser, tmp_path)
    assert first["delivery_ready"] is False
    assert any(item["status"] == "image_bio_email_conflict" for item in first["unresolved"])
    item = collect_email_conflict_items(first["audit_path"])[0]
    save_email_decision(tmp_path / "image_bio_email_reviews.json", item, "234567@163.com")
    second = await run_plan(plan, browser, tmp_path)
    assert second["delivery_ready"] is True
    records = json.loads(Path(second["records_path"]).read_text(encoding="utf-8"))
    assert records[0]["邮箱"] == "234567@163.com"


@pytest.mark.asyncio
async def test_image_cards_reconcile_repeated_text_roster_without_manual_candidates(tmp_path, monkeypatch):
    import server.image_biography as media

    url = "https://faculty.example.edu/mentors"
    names = ("张富有", "刘圣洁", "孙洁", "赵文成")
    listing = "、".join(names)
    html = ('<html><body><article><p>学术型硕导（4人）：' + listing + '</p>'
            '<p>专业型硕导（4人）：' + listing + '</p>'
            + ''.join(f'<img src="/{i}.png">' for i in range(4)) + '</article></body></html>')
    assets = {f"https://faculty.example.edu/{i}.png": str(i).encode() for i in range(4)}
    browser = FakeContext({url: html}, assets=assets)

    def scan(data, _kind, *, ocr_engine=None):
        name = names[int(data)]
        lines = (f"{name} 教授", "完整学习经历与任职信息", "研究项目与学术成果详述")
        text = "\n".join(lines)
        return {"page_count": 1, "pages_ocrd": 1, "text": text,
                "pages": [{"page": 1, "text": text, "source": "local_ocr",
                           "lines": [{"text": line, "confidence": 0.99} for line in lines]}],
                "image_size": [1200, 1800]}

    monkeypatch.setattr(media, "scan_media", scan)
    result = await run_plan({"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "single_page": True, "card_selector": "article img"}]}, browser, tmp_path)

    assert result["delivery_ready"] is True
    assert result["total"] == 4
    assert result["roster_observations"][0]["media_reconciled"] is True
    records = json.loads(Path(result["records_path"]).read_text(encoding="utf-8"))
    assert {row["姓名"] for row in records} == set(names)
    assert all(row["姓名"] in row["简介"] for row in records)


@pytest.mark.asyncio
async def test_shared_scan_pdf_link_on_roster_blocks_unassigned_biographies(tmp_path, monkeypatch):
    import server.image_biography as media

    url = "https://faculty.example.edu/mentors"
    html = ('<html><body><article><ul><li>张富有 教授</li><li>刘圣洁 副教授</li></ul>'
            '<a href="/cv.pdf">导师简介.pdf</a></article></body></html>')
    browser = FakeContext({url: html}, assets={"https://faculty.example.edu/cv.pdf": b"pdf"})
    monkeypatch.setattr(media, "scan_media", lambda *_args, **_kwargs: {
        "page_count": 1, "pages_ocrd": 1, "text": "张富有 教授\n完整导师简介",
        "pages": [{"page": 1, "text": "张富有 教授\n完整导师简介", "source": "local_ocr",
                   "lines": [{"text": "张富有 教授", "confidence": 0.99}]}]})
    result = await run_plan({"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "single_page": True, "card_selector": "article li"}]}, browser, tmp_path)
    assert result["delivery_ready"] is False
    assert result["excel_path"] is None
    assert any(item["status"] == "image_bio_owner_unresolved" for item in result["unresolved"])
    assert len(result["image_biography_evidence"]) == 1


@pytest.mark.asyncio
async def test_partial_poster_review_blocks_export(tmp_path):
    url = "https://faculty.example.edu/poster"
    screenshot = tmp_path / "poster.png"
    screenshot.write_bytes(b"reviewed image")
    plan = {"school": "样本大学", "college": "样本学院", "branches": [
        {"url": url, "single_page": True}
    ], "poster_reviews": [{"source_url": url, "reviewed_full_image": True, "expected_count": 2}],
    "verified_candidates": [
        {"name": "张富有", "title": "教授", "source_url": url, "evidence_path": str(screenshot)},
    ]}

    result = await run_plan(plan, FakeContext({url: "<html><body><img src='poster.png'></body></html>"}), tmp_path)

    assert result["delivery_ready"] is False
    assert result["excel_path"] is None
    assert any(item["status"] == "poster_coverage_mismatch" for item in result["unresolved"])


def test_profileless_candidate_on_second_branch_does_not_trigger_profile_fetch():
    second_roster = "https://faculty.example.edu/department-b"
    candidate = FacultyCandidate(name="张富有", url=second_roster, title="教授")
    called = []

    def fetch(url):
        called.append(url)
        raise AssertionError("directory page is not an individual profile")

    result = FacultyHarvestPipeline.harvest_candidates(
        [candidate], fetch, base_url="https://faculty.example.edu/department-a",
        profileless_urls={second_roster}, school="样本大学", college="样本学院"
    )

    assert called == []
    assert result["delivery_ready"] is True
    assert result["sheet2_no_email"][0]["主页"] == second_roster


def test_pagination_polling_does_not_mutate_cached_wrapper_registry(monkeypatch):
    def fail_if_cache_opened():
        raise AssertionError("transient AJAX DOM must not touch wrapper cache")

    monkeypatch.setattr("server.faculty_extractor.get_default_l1_executor", fail_if_cache_opened)
    signature = _roster_signature(roster([(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]),
                                  "https://faculty.example.edu/roster")
    assert len(signature) == 3


@pytest.mark.asyncio
async def test_real_playwright_browser_handles_ajax_pagination(tmp_path):
    playwright_api = pytest.importorskip("playwright.async_api")
    first = [(1, "张富有"), (2, "刘圣洁"), (3, "孙洁")]
    second = [(4, "赵文成"), (5, "王英华"), (6, "周立新")]

    def cards(people):
        return "".join(f'<li><a href="/p/{i}.htm">{name}</a><span>教授</span></li>'
                       for i, name in people)

    second_cards = cards(second)
    roster_page = (
        '<html><body><main><ul id="people">' + cards(first) + '</ul>'
        '<button class="next">下一页</button>'
        '<script>document.querySelector(".next").addEventListener("click", function () {'
        'document.querySelector("#people").innerHTML='
        + json.dumps(second_cards, ensure_ascii=False) + ';this.disabled=true;});</script>'
        '</main></body></html>'
    )
    profile_pages = {f"/p/{i}.htm": profile(name, f"{i:010d}@example.edu")
                     for i, name in first + second}

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            body = roster_page if self.path == "/roster" else profile_pages.get(self.path)
            if body is None:
                self.send_error(404)
                return
            encoded = body.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(encoded)))
            self.end_headers()
            self.wfile.write(encoded)

        def log_message(self, *_args):
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        async with playwright_api.async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            try:
                context = await browser.new_context()
                url = f"http://127.0.0.1:{server.server_port}/roster"
                plan = {"school": "样本大学", "college": "样本学院", "branches": [
                    {"url": url, "next_selector": ".next", "max_pages": 3,
                     "card_selector": "main li"}
                ]}
                result = await run_plan(plan, context, tmp_path)
                await context.close()
            finally:
                await browser.close()
    finally:
        server.shutdown()
        thread.join(timeout=5)
        server.server_close()

    assert result["delivery_ready"] is True
    assert result["directory_pages"] == 2
    assert result["profile_pages"] == 6
    assert result["total"] == 6
