"""Full native text survives workbook, sync, leased persistence and publishing."""
import copy
import json

import openpyxl
import pytest

from server.agent_queue import AgentTaskQueue
from server.browser_batch_sync import build_payload
from server.excel_export import export_dual_sheet_excel, verify_biography_storage
from server.university_delivery import publish_delivery
from test_delivery_field_receipt import claim, record
from test_browser_batch_sync import _report


def reconstructed(path):
    workbook = openpyxl.load_workbook(path)
    try:
        return ''.join(row[9] for row in list(workbook['完整简介'].values)[1:])
    finally:
        workbook.close()


def test_unicode_linebreaks_and_formula_like_native_text_are_lossless(tmp_path):
    original = '=完整简介\n' + '中文😀\n' * 12000 + '  末尾\n'
    rows = [record(简介=original)]
    before = copy.deepcopy(rows)
    path = tmp_path / 'full.xlsx'
    counts = export_dual_sheet_excel(rows, str(path))
    assert rows == before
    assert reconstructed(path) == original
    assert verify_biography_storage(path, rows) == counts['biography_storage']
    book = openpyxl.load_workbook(path)
    try:
        assert all(len(row[9].encode('utf-16-le')) // 2 <= 32767
                   for row in list(book['完整简介'].values)[1:])
        assert book['完整简介']['J2'].data_type == 's'
        assert book['无邮箱教师列表']['I2'].hyperlink.location == "'完整简介'!A2"
    finally:
        book.close()


@pytest.mark.parametrize('damage', ['fragment', 'owner', 'link', 'missing'])
def test_written_artifact_damage_is_rejected(tmp_path, damage):
    rows = [record(简介='甲' * 40000)]
    path = tmp_path / 'full.xlsx'
    export_dual_sheet_excel(rows, str(path))
    book = openpyxl.load_workbook(path)
    if damage == 'fragment':
        book['完整简介']['J2'] = '截断'
    elif damage == 'owner':
        book['完整简介']['D2'] = '另一个人'
    elif damage == 'link':
        book['无邮箱教师列表']['I2'].hyperlink.location = "'完整简介'!A3"
    else:
        del book['完整简介']
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match='full biography'):
        verify_biography_storage(path, rows)


def test_distinct_people_keep_original_partition_and_chunk_ownership(tmp_path):
    rows = [record(简介='甲' * 70000),
            record(姓名='第二人', 邮箱='second@example.edu', 主页='https://example.edu/second',
                   简介='乙' * 40000)]
    path = tmp_path / 'full.xlsx'
    counts = export_dual_sheet_excel(rows, str(path))
    assert (counts['total'], counts['with_email'], counts['no_email']) == (2, 1, 1)
    assert verify_biography_storage(path, rows) == counts['biography_storage']
    book = openpyxl.load_workbook(path)
    try:
        assert book['教师列表']['B2'].value == '第二人'
        assert book['无邮箱教师列表']['B2'].value == '合成人名'
        assert book['教师列表']['I2'].hyperlink.location == "'完整简介'!A5"
    finally:
        book.close()


def test_record_bound_receipt_is_required_before_leased_database_commit(tmp_path):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task, leased = claim(queue)
    rows = [record(简介='甲' * 40000)]
    counts = export_dual_sheet_excel(rows, str(tmp_path / 'full.xlsx'))
    tampered = copy.deepcopy(counts)
    tampered['biography_storage']['receipts'][0]['sha256'] = '0' * 64
    with pytest.raises(ValueError, match='storage receipt'):
        queue.complete_task(task.task_id, rows, tampered, worker_id='worker',
                            lease_token=leased.lease_token)
    assert queue.get_task(task.task_id).records == []
    completed = queue.complete_task(task.task_id, rows, counts, worker_id='worker',
                                    lease_token=leased.lease_token)
    assert completed.records == rows
    delivery = publish_delivery(queue, task.batch_id, task.school, tmp_path / 'outputs')
    assert delivery['total'] == 1
    assert reconstructed(delivery['excel_path']) == rows[0]['简介']


def test_sync_reads_actual_full_workbook_and_refuses_forged_receipt(tmp_path):
    path, root = _report(tmp_path)
    report = json.loads(path.read_text(encoding='utf8'))
    records_path = tmp_path / 'records.json'
    rows = json.loads(records_path.read_text(encoding='utf8'))
    rows[0]['简介'] = '完整履历\n' * 10000
    records_path.write_text(json.dumps(rows, ensure_ascii=False), encoding='utf8')
    counts = export_dual_sheet_excel(rows, report['excel_path'])
    report['biography_storage'] = counts['biography_storage']
    path.write_text(json.dumps(report, ensure_ascii=False), encoding='utf8')
    payload = build_payload(path, 'worker', 'lease', root)
    assert payload['records'] == rows
    assert payload['stats']['biography_storage'] == counts['biography_storage']
    report['biography_storage']['receipts'][0]['parts'] += 1
    path.write_text(json.dumps(report, ensure_ascii=False), encoding='utf8')
    with pytest.raises(ValueError, match='receipt differs'):
        build_payload(path, 'worker', 'lease', root)
