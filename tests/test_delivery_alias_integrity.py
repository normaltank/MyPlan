"""Legacy receipts stay immutable while their university workbook uses ten columns."""
import json
import openpyxl
import pytest

from server.agent_queue import AgentTaskQueue
from server.university_delivery import publish_delivery


@pytest.mark.parametrize('scope', [{}, {'school': '甲大学', 'college': '甲学院'}])
def test_legacy_receipt_publishes_known_scope_and_exact_biography(tmp_path, scope):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task = queue.submit_task('甲大学', '甲学院')
    leased = queue.claim_task('worker', task_id=task.task_id)
    raw = {'name': '王芳', 'email': 'wang@example.edu', 'title': '教授',
           'full_bio': '  原始简介\n\n完整履历。  \n', **scope}
    queue.complete_task(task.task_id, [raw], {'total': 1, 'with_email': 1, 'no_email': 0},
                        worker_id='worker', lease_token=leased.lease_token)
    delivery = publish_delivery(queue, task.batch_id, task.school, tmp_path / 'outputs')
    assert queue.get_task(task.task_id).records == [raw]
    workbook = openpyxl.load_workbook(delivery['excel_path'])
    row = list(workbook['教师列表'].values)[1]
    assert row[1:6] == ('王芳', '教授', 'wang@example.edu', '甲大学', '甲学院')
    assert row[8] == raw['full_bio']
    assert (delivery['total'], delivery['with_email'], delivery['no_email']) == (1, 1, 0)
    workbook.close()


def test_empty_canonical_email_uses_nonempty_legacy_value_consistently(tmp_path):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task = queue.submit_task('甲大学', '甲学院')
    leased = queue.claim_task('worker', task_id=task.task_id)
    raw = {'姓名': '王芳', '邮箱': '', 'email': 'wang@example.edu',
           '学校': '甲大学', '学院': '甲学院', '简介': '完整简介'}
    queue.complete_task(task.task_id, [raw], {'total': 1, 'with_email': 1, 'no_email': 0},
                        worker_id='worker', lease_token=leased.lease_token)
    delivery = publish_delivery(queue, task.batch_id, task.school, tmp_path / 'outputs')
    assert queue.get_task(task.task_id).records == [raw]
    workbook = openpyxl.load_workbook(delivery['excel_path'])
    assert workbook['教师列表'].max_row == 2 and workbook['无邮箱教师列表'].max_row == 1
    assert workbook['教师列表']['D2'].value == 'wang@example.edu'
    assert delivery['with_email'] == 1
    workbook.close()


@pytest.mark.parametrize('fields', [
    {'姓名': '王芳', 'name': '李华'},
    {'邮箱': 'wang@example.edu', 'email': 'li@example.edu'},
    {'职称': '教授', 'title': '讲师'},
    {'简介': '甲履历', 'full_bio': '乙履历'},
    {'主页': 'https://example.edu/a', 'profile_url': 'https://example.edu/b'},
])
def test_conflicting_nonempty_field_aliases_cannot_complete(tmp_path, fields):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task = queue.submit_task('甲大学', '甲学院')
    leased = queue.claim_task('worker', task_id=task.task_id)
    with_email = int(bool(fields.get('email') or fields.get('邮箱')))
    with pytest.raises(ValueError, match='record_field_alias_conflict'):
        queue.complete_task(task.task_id, [{'姓名': '王芳', **fields}],
                            {'total': 1, 'with_email': with_email, 'no_email': 1 - with_email},
                            worker_id='worker', lease_token=leased.lease_token)
    saved = queue.get_task(task.task_id)
    assert saved.status == 'running' and saved.records == []


def test_scope_less_reviewed_homonyms_keep_their_original_identity_certificate(tmp_path):
    from pathlib import Path
    from server import identity_review
    from server.university_delivery import manifest_snapshot_is_safe
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    task = queue.submit_task('甲大学', '甲学院')
    leased = queue.claim_task('worker', task_id=task.task_id)
    observations = [{'name': '王芳', 'profile_url': f'https://example.edu/{index}',
                     'email': f'person{index}@example.edu', 'full_bio': f'  完整履历{index}  '}
                    for index in (1, 2)]
    hashes = {row['profile_url']: 'a' * 64 for row in observations}
    evidence = tmp_path / 'tasks' / task.task_id / 'attempt-1'
    evidence.mkdir(parents=True)
    audit = evidence / 'audit.json'
    audit.write_text('{}', encoding='utf-8')
    path = evidence / identity_review.DECISIONS_FILE
    item = identity_review.pending_items(observations, hashes)[0]
    identity_review.save_decision(path, item, 'distinct', reviewed_by=1)
    records, pending, certificates = identity_review.apply_decisions(observations, hashes, path)
    assert not pending
    stats = {'total': 2, 'with_email': 2, 'no_email': 0, 'identity_reviews': certificates,
             'identity_review_path': str(path), 'audit_path': str(audit)}
    queue.complete_task(task.task_id, records, stats, worker_id='worker', lease_token=leased.lease_token)
    delivery = publish_delivery(queue, task.batch_id, task.school, tmp_path / 'outputs')
    assert delivery['delivery_scope']['state'] == 'ready' and delivery['total'] == 2
    assert queue.get_task(task.task_id).records == records
    manifest = json.loads(Path(delivery['manifest_path']).read_text(encoding='utf-8'))
    assert manifest_snapshot_is_safe(queue, manifest)


def test_complementary_observation_copies_the_full_raw_biography(tmp_path):
    queue = AgentTaskQueue(tmp_path / 'queue.db')
    bio = '  原始简介\n\n完整履历。  \n'
    for college, biography in [('甲学院', ''), ('乙学院', bio)]:
        task = queue.submit_task('甲大学', college, batch_id='batch')
        leased = queue.claim_task('worker', task_id=task.task_id)
        raw = {'name': '王芳', 'profile_url': 'https://example.edu/profile', 'full_bio': biography}
        queue.complete_task(task.task_id, [raw], {'total': 1, 'with_email': 0, 'no_email': 1},
                            worker_id='worker', lease_token=leased.lease_token)
    delivery = publish_delivery(queue, 'batch', '甲大学', tmp_path / 'outputs')
    workbook = openpyxl.load_workbook(delivery['excel_path'])
    assert workbook['无邮箱教师列表']['I2'].value == bio
    assert delivery['total'] == 1
    workbook.close()
