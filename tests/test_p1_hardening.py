"""P1 hardening: native Chrome isolation, no site-specific WAF bypass,
fail-fast worker self-check and text-only Excel cells."""

import concurrent.futures
import sys
from pathlib import Path

import openpyxl
import pytest

import server.browser_batch_tool as browser_batch_tool
from server.autonomous_agent import AutonomousBrowserAgent


def _record(name, bio, email=""):
    return {"序号": 1, "姓名": name, "职称": "教授", "邮箱": email, "学校": "U", "学院": "C",
            "采集人": "fixture", "采集日期": "2026-09-29", "简介": bio, "主页": f"https://u.edu/{name}"}


# ── Excel cells are text ─────────────────────────────────────────────────────
def test_harvested_text_is_never_stored_as_a_formula(tmp_path):
    path = tmp_path / "u.xlsx"
    AutonomousBrowserAgent().export_dual_sheet_excel(
        [_record("甲", '=HYPERLINK("http://evil.example","点我")'),
         _record("乙", "=1+1", "b@u.edu")], str(path))
    cells = [row[8] for ws in openpyxl.load_workbook(path).worksheets
             for row in ws.iter_rows(min_row=2)]
    assert {cell.value for cell in cells} == {'=HYPERLINK("http://evil.example","点我")', "=1+1"}
    assert all(cell.data_type == "s" for cell in cells)


def test_control_characters_no_longer_abort_the_university_workbook(tmp_path):
    path = tmp_path / "u.xlsx"
    AutonomousBrowserAgent().export_dual_sheet_excel(
        [_record("乙", "履历\x0b第二段\x01完整")], str(path))
    bio = openpyxl.load_workbook(path)["无邮箱教师列表"]["I2"].value
    assert bio == "履历第二段完整"


# ── No site-specific WAF bypass in the general tool ─────────────────────────
def test_general_batch_tool_has_no_site_specific_waf_bypass():
    source = Path(browser_batch_tool.__file__).read_text(encoding="utf-8")
    assert not hasattr(browser_batch_tool, "_setup_waf_proxy_if_needed")
    assert "gszy.edu.cn" not in source
    assert "verify=False" not in source


# ── Native Chrome fallback isolation ────────────────────────────────────────
class _FakeChrome:
    launched = []

    def __init__(self, args, **_kwargs):
        self.args = args
        user_dir = Path(next(a for a in args if a.startswith("--user-data-dir=")).split("=", 1)[1])
        port = 40000 + len(_FakeChrome.launched)
        (user_dir / "DevToolsActivePort").write_text(f"{port}\n/devtools/browser/fake\n", encoding="ascii")
        _FakeChrome.launched.append((user_dir, port))
        self.stopped = False

    def poll(self):
        return None

    def terminate(self):
        self.stopped = True

    def wait(self, timeout=None):
        return 0


def test_each_native_chrome_gets_its_own_port_and_profile_which_is_removed(monkeypatch, tmp_path):
    import subprocess
    _FakeChrome.launched = []
    monkeypatch.setattr(subprocess, "Popen", _FakeChrome)
    first = browser_batch_tool._launch_native_chrome(tmp_path / "chrome", ["https://a.edu/"])
    second = browser_batch_tool._launch_native_chrome(tmp_path / "chrome", ["https://b.edu/"])
    assert "--remote-debugging-port=0" in first[0].args
    assert first[2] != second[2] and first[1] != second[1]
    for proc, user_dir, _port in (first, second):
        browser_batch_tool._stop_native_chrome(proc, user_dir)
        assert proc.stopped and not user_dir.exists()


def test_concurrent_native_chrome_fetches_do_not_cross(monkeypatch, tmp_path):
    monkeypatch.delenv("HARVEST_DISABLE_NATIVE_CHROME_FALLBACK", raising=False)
    if browser_batch_tool._native_chrome_executable() is None:
        pytest.skip("no local Chrome/Chromium for the native fallback")
    pages = []
    for index in range(3):
        page = tmp_path / f"college{index}.html"
        page.write_text(f"<html><body><h1>college-{index}</h1></body></html>", encoding="utf-8")
        pages.append(page.as_uri())
    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
        results = list(pool.map(browser_batch_tool._fetch_page_via_native_chrome, pages))
    for index, html in enumerate(results):
        assert html is not None and f"college-{index}" in html


# ── Fail-fast worker self-check ─────────────────────────────────────────────
def test_self_check_reports_each_problem_and_survives_a_broken_check():
    from server.worker_selfcheck import run_self_check

    def broken():
        raise RuntimeError("boom")
    problems = run_self_check((("ok", lambda: None), ("missing", lambda: "缺少组件"), ("broken", broken)))
    assert problems == ["[missing] 缺少组件", "[broken] 自检异常：RuntimeError: boom"]


def test_algorithm_layer_and_pinyin_checks_pass_in_a_complete_environment():
    from server import worker_selfcheck
    assert worker_selfcheck._l2_algorithms() is None
    assert worker_selfcheck._pinyin() is None


def test_missing_api_key_is_reported(monkeypatch):
    from server import worker_selfcheck
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    assert "DEEPSEEK_API_KEY" in worker_selfcheck._api_key()


def test_worker_refuses_to_start_when_self_check_fails(monkeypatch, capsys):
    from server import agent_worker, worker_selfcheck
    monkeypatch.setattr(worker_selfcheck, "run_self_check", lambda: ["[拼音] 缺少 pypinyin"])
    monkeypatch.setattr(sys, "argv", ["agent_worker"])
    constructed = []
    monkeypatch.setattr(agent_worker, "AgentWorker", lambda **kw: constructed.append(kw))
    with pytest.raises(SystemExit) as stopped:
        agent_worker.main()
    assert stopped.value.code == 2 and constructed == []
    assert "缺少 pypinyin" in capsys.readouterr().err
