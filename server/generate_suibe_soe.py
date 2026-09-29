import os
import sys
import json
import urllib.request
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

def main():
    school = "上海对外经贸大学"
    college = "高级研修学院"
    gatherer = "Antigravity Subagent (Browser)"
    gather_date = "2026-09-21"
    task_id = "08e83d3a"

    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    
    # 0 teachers found as 高级研修学院 is a continuing education / executive development administrative platform with 0 dedicated faculty
    records = []
    
    df = pd.DataFrame(records, columns=columns)
    df_with_email = df[df["邮箱"].notna() & (df["邮箱"] != "")].copy()
    df_no_email = df[df["邮箱"].isna() | (df["邮箱"] == "")].copy()

    sheet1_name = f"有邮箱学者 (共 {len(df_with_email)} 位)"
    sheet2_name = f"无邮箱学者 (共 {len(df_no_email)} 位)"

    os.makedirs("server/outputs", exist_ok=True)
    timestamp = pd.Timestamp.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{school}_{college}_{timestamp}.xlsx"
    filepath = os.path.join("server/outputs", filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df_with_email.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_no_email.to_excel(writer, sheet_name=sheet2_name, index=False)

    # Style with Openpyxl
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    data_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            if row_idx == 1:
                ws.row_dimensions[row_idx].height = 28
                for cell in row:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border
            else:
                ws.row_dimensions[row_idx].height = 22
                for cell in row:
                    cell.font = data_font
                    cell.alignment = Alignment(vertical="center")
                    cell.border = thin_border
        
        # Set column widths
        col_widths = {
            'A': 8,   # 序号
            'B': 14,  # 姓名
            'C': 20,  # 职称
            'D': 25,  # 邮箱
            'E': 18,  # 学校
            'F': 18,  # 学院
            'G': 28,  # 采集人
            'H': 14,  # 采集日期
            'I': 50,  # 简介
            'J': 45   # 主页
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

    wb.save(filepath)
    print(f"✅ Excel 已生成: {filepath}")

    # Synchronize with backend
    sync_url = f"http://127.0.0.1:8000/api/task/{task_id}/sync"
    sync_data = {
        "records": records,
        "stats": {
            "total": 0,
            "with_email": 0,
            "no_email": 0
        },
        "download_url": f"/download/{filename}",
        "progress": 100
    }

    req = urllib.request.Request(
        sync_url,
        data=json.dumps(sync_data).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    
    try:
        resp = urllib.request.urlopen(req, timeout=10)
        res_json = json.loads(resp.read().decode("utf-8"))
        print(f"✅ 后台同步成功: {res_json.get('status')}")
    except Exception as e:
        print(f"❌ 同步失败: {e}")

if __name__ == "__main__":
    main()
