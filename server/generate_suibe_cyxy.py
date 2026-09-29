import os
import sys
import json
import urllib.request
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

def main():
    task_id = "1d34b2fc"
    school = "上海对外经贸大学"
    college = "创业学院"
    gatherer = "Antigravity Subagent (Browser)"
    gather_date = "2026-09-21"

    teachers = [
        {
            "name": "沈艺",
            "title": "副研究员",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2026/0107/c12271a192292/page.htm",
            "intro": "沈艺，女，1978年9月出生，汉族，上海人，研究生学历，管理学博士，副研究员，现任上海对外经贸大学党委副书记、副校长。1999年加入中国共产党，2001年7月参加工作，曾担任上海大学党委统战部副部长、校党委组织部副部长，社会学院党委书记，校党委常委、组织人事部常务副部长、组织处处长，党校常务副校长、机构编制委员会办公室主任等职务。曾主持、参与完成多项省部级研究课题，全国党建研究重点课题等，发表出版研究论文和专著多篇（部），兼任上海市行政管理学会常务理事。2025年9月，兼任上海对外经贸大学创业学院院长。"
        },
        {
            "name": "高伟",
            "title": "教授",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2022/0409/c12271a145172/page.htm",
            "intro": "高伟，博士，上海对外经贸大学创业学院（中国—上海合作组织经贸学院创新创业学院）副院长、研究生导师，美国管理技术大学特聘教授，中国—上海合作组织经贸学院荣誉教授，可持续发展Mover社群助力官、上海市阳光学者、联合国教科文组织中国创业教育联盟常务理事、中国创造学会新质生产力产教融合分会常务副主任委员兼秘书长、上海安徽经济文化促进会淮北分会常务副会长兼青才联主席、上海市创业指导专家志愿团理事会副理事长、中国青年创业导师、人力资源社会保障部创业培训讲师、上海市东方讲坛及前程无忧特聘讲师，国际劳工组织KAB培训师、教育部本科教育教学审核评估专家、中国高校创新创业教育研究中心首批专家、工信部中小企业发展促进中心首批产教融合产业实践教授、全国生涯规划专家、高级职业指导师、高级企业培训师、BCC全球生涯教练、牛津大学展望计划访问学者，全国百强创业社团——开源创新创业创造社指导老师，国创赛、大小挑、电商三创赛、职业院校创赛等资深评审；主持教育部高校辅导员工作精品项目、教育部产教融合协同育人项目、教育部供需对接就业育人项目、中国科协学风涵养工作室、上海高校学生职业生涯发展教育示范工作室、上海高校一流本科课程、上海市创业培训精品课程、kab创新实验室、中国—上海合作组织青年创业国际孵化器（上海）专家工作室，主持开设创业管理微专业、“创新创业创造+融合教育”在线微专业，曾获全国高校辅导员年度人物提名奖、全国首届创业指导课程教学大赛优秀奖、全国商业科技进步三等奖、上海市教学成果特等奖、上海市育才奖、上海青年五四奖章、上海高校辅导员年度人物、深化高校创新创业教育改革优秀教材、著作、实践及论文等荣誉。在解放日报理论版、《思想理论教育》、《Entrepreneurship Education》、《国际商务研究》等核心期刊等发表中英文学术论文20余篇，出版学术著作9部，参与及完成国家社科重点项目、教育部人文社科青年项目等省级以上科研课题20余项。参与指导学生团队荣获中国国际大学生创新大赛国赛金奖多项、银奖两项、大挑国赛一等奖、小挑国赛金奖、中国机器人及人工智能国赛一等奖、全国高校商业精英挑战赛国赛一等奖等荣誉。"
        },
        {
            "name": "李伟铭",
            "title": "讲师",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2022/0410/c18644a145230/page.htm",
            "intro": "创新创业项目孵化运营中心负责人，上海市大学生科技创业基金会上海对外经贸大学分会负责人，工信部中小企业志愿服务专家，中国国际“互联网+”大学生创新创业大赛国赛评委，KAB创业教育（中国）项目讲师，新媒体营销师，国际中级创业教练，《创新创业创造：职场竞争力密钥》慕课教研组成员；长期致力于创新创业课程教学、创新创业竞赛指导和创业项目孵化，讲授“大学生KAB创业基础”、“创业管理”、“企业战略管理”、“市场营销”等多门课程，指导多个大学生创业项目参加“互联网+”、“挑战杯”、“创青春”、“中国创翼”、“退役军人创业创新大赛”、“女大学生创业大赛”等各级各类创新创业竞赛，荣获“互联网+”大赛全国总决赛铜奖和上海市金奖，获评“互联网+”大赛市级优秀指导教师；运营管理上海对外经贸大学“文汇141”创业咖啡、“文汇580”大学生创业中心和“古北620”创业孵化基地近3000平米创业项目孵化空间，指导超60个创业项目落地开业。"
        },
        {
            "name": "李朝阳",
            "title": "讲师",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2022/0410/c18644a145231/page.htm",
            "intro": "在上海对外经贸大学从事创业教育教学、研究和教学管理相关工作，自2017底在新成立的创业学院从事创业教育相关工作。具有10年多高校创业类课程授课经验，具有国家人社部颁发的开业咨询师（二级）证、KAB创业受训证、高校教师资格证，著有《融入商科人才培养体系的创业实训课程的建设》、《大学生创业基础（第2版）》、《创业教育、思政教育和专业教育融合的实践探索》等研究成果。"
        },
        {
            "name": "张博",
            "title": "教学秘书",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2022/0411/c18644a145275/page.htm",
            "intro": "教学秘书，国际创业教练（中级），国际职业导师（中级），“职业生涯设计”模块课程 教学管理，《高校创新创业教育前沿》内参编辑，指导学生参加暑期社会实践，获校一等奖，指导多个学生团队参与大学生创新创业训练计划获市级立项。"
        },
        {
            "name": "施毅婷",
            "title": "副教授",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2022/0410/c18645a145235/page.htm",
            "intro": "会计学院党委副书记，指导《“烟火匠的信仰”公益团队——以协同服务赋能基层治理的一站式公益方案提供者》荣获第七届中国国际“互联网+”大学生创新创业大赛全国总决赛铜奖和上海市金奖，并获评优秀指导教师，指导多个大学生创业项目在“互联网+”大赛、“挑战杯”等创新创业赛事获奖，是《2020胡润Under30s创业领袖》、“独立日”创始人郑一创业启蒙导师。曾获 2016 年上海市育才奖， 上海市教育系统“三八红旗手”称号，全国第五届辅导员素质能力大赛二等奖、 华东赛区一等奖，负责上海市教委职业生涯规划工作室一项，2021 年上海市高 校教师就业指导课程教学大赛三等奖，2018 年度上海市辅导员“年度人物”提名奖，2015 年辅导员素质能力大赛上海赛区二等奖，2015 年上海市大学生暑期社会实践活动优秀指导教师。"
        },
        {
            "name": "于卓",
            "title": "讲师",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2022/0411/c18645a145248/page.htm",
            "intro": "国际经贸学院辅导员、上海高校学生职业（生涯）发展教育示范工作室（“3F-职业志”生涯育人工作室）主创成员、教育部首期全国职业教育“新媒体营销专项”师资认证、EEGC高效能就业团体辅导认证（中级）、SFBT焦点解决生涯辅导技术认证（中级）、BCC国际生涯教练、全国商业精英挑战赛国际贸易竞赛 优秀指导老师、上海市商业精英挑战赛国际贸易竞赛 优秀指导老师、长期指导校内学生参加“互联网+”、“挑战杯”、“创青春”等创业赛事。"
        },
        {
            "name": "徐同飞",
            "title": "专职辅导员",
            "email": "",
            "homepage": "https://www.suibe.edu.cn/cyxy/2022/0411/c18645a145252/page.htm",
            "intro": "徐同飞，上海对外经贸大学金融管理学院专职辅导员，上海市青年研究会会员，主持上海共青团相关课题两项，发表创新创业相关论文多篇，原彩虹教育、树人总动员等商业和公益创业团队负责人。曾主持研究生项目获得“挑战杯”大学生创业计划竞赛国家级奖项；指导本科生团队获得中国国际“互联网+”大学生创新创业大赛国家级奖项，取得学校历史上零的突破。指导学生获“互联网+”大学生创新创业大赛上海市银奖两项等，指导国家级和上海市级大学生创新创业训练计划多项。持高等教育教师资格证，参与形势与政策、职业生涯发展与规划、KAB大学生创业基础等课程授课。"
        }
    ]

    records = []
    for idx, t in enumerate(teachers, 1):
        records.append({
            "序号": idx,
            "姓名": t["name"],
            "职称": t["title"],
            "邮箱": t["email"] or None,
            "学校": school,
            "学院": college,
            "采集人": gatherer,
            "采集日期": gather_date,
            "简介": t["intro"],
            "主页": t["homepage"]
        })

    df = pd.DataFrame(records)
    df_with_email = df[df["邮箱"].notna() & (df["邮箱"] != "")].copy()
    df_no_email = df[df["邮箱"].isna() | (df["邮箱"] == "")].copy()

    df_with_email["序号"] = range(1, len(df_with_email) + 1)
    df_no_email["序号"] = range(1, len(df_no_email) + 1)

    sheet1_name = f"有邮箱学者 (共 {len(df_with_email)} 位)"
    sheet2_name = f"无邮箱学者 (共 {len(df_no_email)} 位)"

    os.makedirs("server/outputs", exist_ok=True)
    timestamp = pd.Timestamp.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{school}_{college}_{timestamp}.xlsx"
    filepath = os.path.join("server/outputs", filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df_with_email.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_no_email.to_excel(writer, sheet_name=sheet2_name, index=False)

    # Openpyxl Styling
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    data_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = data_font
                    if cell.column in [1, 2, 3, 5, 6, 7, 8]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col[:15]:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str.encode('gbk', 'ignore')))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    wb.save(filepath)
    print(f"🎉 Standard Dual-Sheet Excel saved to: {filepath}")
    print(f"  {sheet1_name}")
    print(f"  {sheet2_name}")
    print(f"  Total teachers: {len(df)} (Strictly Conserved: {len(df_with_email)} + {len(df_no_email)} = {len(df)})")

    # Sync to backend
    sync_url = f"http://127.0.0.1:8000/api/task/{task_id}/sync"
    sync_payload = {
        "records": records,
        "stats": {
            "total": len(records),
            "with_email": len(df_with_email),
            "no_email": len(df_no_email),
            "email_rate": round(len(df_with_email) / len(records) * 100, 1) if records else 0
        },
        "download_url": f"/download/{filename}",
        "progress": 100
    }

    req = urllib.request.Request(
        sync_url,
        data=json.dumps(sync_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )

    try:
        with urllib.request.urlopen(req) as resp:
            resp_data = resp.read().decode("utf-8")
            print(f"✅ Sync API Response: {resp_data}")
    except Exception as e:
        print(f"❌ Sync API Error: {e}")

if __name__ == "__main__":
    main()
