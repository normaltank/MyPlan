import os
import sys
import io
import json
import re
import datetime
import urllib.request
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

FACULTY_DATA = [
    {
        "姓名": "任丽杰",
        "职称": "教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "任丽杰，教授，心理学博士，上海海关学院马克思主义学院副院长，党支部书记，上海市心理学会青春期与性心理工作委员会副主任，上海高校心理咨询协会常务理事，北京生命教育促进会生命教育工作室负责人，中国心理学会注册心理师，中国心理卫生协会注册心理师，上海学校心理咨询师（高级）。主要从事大学生心理健康和生命教育的教学和研究，在国内外学术期刊发表学术论文40余篇，出版专著2部，主讲的课程《大学生生命教育》《大学生心理健康》分别获得上海市高校教师教学创新大赛二等奖。曾获全国心理健康教育优秀工作者，上海市育才奖，上海学校心理抗疫先进个人，上海高校心理健康教育先进个人等荣誉称号。",
        "主页": "https://www.shcc.edu.cn/2026/0724/c1298a70092/page.htm"
    },
    {
        "姓名": "钱薇",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "钱薇，副教授，硕士，上海海关学院马克思主义学院院长、党支部书记。曾任上海海关学院图书馆馆长、科研党支部书记，曾获评校级优秀党务工作者。主要从事高校思想政治教育、党建引领育人、文献信息资源管理与海关特色智库文献服务等领域的研究与管理工作。主持并推进马克思主义学院思政课改革创新、“大思政课”实践基地建设与学科团队协同创新。",
        "主页": "https://mks.shcc.edu.cn/1587/list.htm"
    },
    {
        "姓名": "姚永超",
        "职称": "教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "姚永超，上海海关学院海关史研究院教授、副院长，曾任马克思主义学院临时负责人。复旦大学海关史与海关文献国际研究中心、厦门大学中国海关史研究中心兼职研究员，英国伦敦大学皇家霍洛威学院历史系访问学者。国家社科基金中国历史研究院重大历史问题研究项目“海内外江海关档案资料整理与研究（1854-1949）”首席专家。近年来，曾主持完成国家社科基金一般项目“中国旧海关海图的整理与研究”，国家社科基金重大项目“中国旧海关内部出版物的整理与研究”与“大阪产业部近代中国及‘海上丝路’沿线调查资料整理与研究”子课题等多项研究任务。在《历史地理》《中国历史地理论丛》《史林》《史学月刊》等核心期刊上发表学术论文40余篇。专著《东北近代经济地理》，作为“中国近代经济地理”大型丛书九卷本之一，荣获教育部第八届和上海市哲学社会科学第十四届优秀学术著作类成果一等奖。译著《潮来潮去：海关与中国现代性的全球起源》，被《中华读书报》等评选为年度最佳历史图书。主讲的《中国海关史》课程，先后被评选为上海市高校重点课程（2019年）、与党史教育相融合的百门示范课程（2021年）、一流课程（2024年）。所著《中国海关史十六讲》（修订版），荣获海关总署海关高等教育精品教材（2025年）。",
        "主页": "https://www.shcc.edu.cn/2019/0621/c1298a30786/page.htm"
    },
    {
        "姓名": "杨寄荣",
        "职称": "教授",
        "邮箱": "jryang16@126.com",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "杨寄荣，复旦大学马克思主义基本原理专业博士。上海海关学院马克思主义学院教授。中国伦理学会理事。上海海关学院学术委员会委员。上海市伦理学会常务理事；上海市哲学学会理事。上海市《毛泽东思想和中国特色社会主义理论体系概论》课教指委委员。2016年获上海市育才奖。2017年获上海市马克思主义理论研究“中青年拔尖人才”称号。发表论文30余篇，著作《发展的伦理进路—基于马克思主义立场》（上海三联书店），《全球化的发展问题》（天津人民出版社），主编教材《海关职业文化》、《思政理论课案例研讨辅学教程》等。主持教育部人文社会科学研究一般项目“新时代民族精神的教学转化研究”等多项省部级科研项目。",
        "主页": "https://www.shcc.edu.cn/2019/0621/c1298a30789/page.htm"
    },
    {
        "姓名": "杨敬敏",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "杨敬敏，上海海关学院马克思主义学院副教授，马克思主义中国化教研室主任，历史学博士。2009年本科毕业于北京师范大学历史系；随后在复旦大学中国历史地理研究所硕博连读，先后师从樊如森教授和吴松弟教授，获得历史学博士学位。主要从事中国经济史、海关史的研究与教学工作。入选上海市“曙光计划”、2025年上海市马克思主义理论学科发展支持计划（马克思主义理论中青年骨干）。著有专著《中国近代棉纺织进口替代工业研究(1867-1936)》（齐鲁书社）。主持国家社科基金年度项目“近代日本涉华海关贸易统计资料整理与研究”、上海市“晨光计划”项目“中国近代保税制度研究”等课题，深度参与国家社科基金重大项目“中国旧海关出版物的整理与研究”及“海内外江海关档案资料整理与研究（1854—1949）”。",
        "主页": "https://mks.shcc.edu.cn/2026/0408/c1590a68761/page.htm"
    },
    {
        "姓名": "顾荣",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "顾荣，博士，上海海关学院马克思主义学院副教授，思想政治教育教研室主任，党支部委员会委员。主要研究方向为马克思主义中国化、思想政治教育理论与实践、高校党建育人机制。承担《马克思主义基本原理概论》、《思想道德修养与法治》、《形势与政策》及《国家安全教育》等核心思政课程教学与集体备课主持工作。主持与参与多项校级和上海市教育科学研究课题，多次在上海高校思政理论课教学观摩与研讨活动中主讲交流，荣获校级党员示范岗等表彰。",
        "主页": "https://mks.shcc.edu.cn/2019/1223/c1653a41344/page.htm"
    },
    {
        "姓名": "吴锋",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "吴锋，上海海关学院马克思主义学院副教授，三十年教龄资深思政教师，曾任思想政治教育教研室主任。深耕高校思想政治理论课一线教学三十载，长期承担《毛泽东思想和中国特色社会主义理论体系概论》、《形势与政策》、《军事理论》等课程的教学组织与实施。荣获上海海关学院“三十年教龄教师风采”专题表彰。积极推进思政理论课教学改革与跨校联合集体备课，深入参与学院本科教学质量保障体系建设。",
        "主页": "https://www.shcc.edu.cn/rsc/jsfz/index.htm"
    },
    {
        "姓名": "潘婷",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "潘婷，女，1987年生，湖南常德人，历史学博士。2019年毕业于上海大学中国近现代史专业，获历史学博士学位。现任上海海关学院马克思主义学院讲师、专任教师。主要研究方向包括上海城市史、中共党史与中国近现代史基本问题。承担《思想道德修养与法治》、《中国近现代史纲要》、《毛泽东思想和中国特色社会主义理论体系概论》等课程教学。在国内外学术期刊发表多篇论文，主持并参与多项市级与校级科研及教学改革项目。",
        "主页": "https://mks.shcc.edu.cn/2026/0909/c1592a70238/page.htm"
    },
    {
        "姓名": "董强",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "董强，历史学博士，讲师。毕业于苏州大学社会学院，获历史学博士学位。现任上海海关学院海关史研究院/马克思主义学院专任教师，中国近代史基本问题学科团队骨干。主要研究领域为中国近现代社会经济史、中国海关史与灾害史。主持国家社科基金青年项目“近代长三角口岸城市灾害协同治理研究”、上海市哲学社会科学规划青年项目等课题。出版学术专著《近代江南公共危机与社会应对》。在《历史地理研究》等核心期刊发表学术论文多篇。承担校级本科课程思政培育项目《文物里的中国》，教学案例《海关是国门生物安全的守护者》入选上海高校思想政治理论课精彩案例。",
        "主页": "https://mks.shcc.edu.cn/2019/1217/c1590a41196/page.htm"
    },
    {
        "姓名": "余越",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "马克思主义学院",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "余越，硕士，现任上海海关学院马克思主义学院党支部委员会委员、学院办公室教学科研秘书、思政课教师。主要负责马克思主义学院党建组织工作、日常行政与教学教务运行、实践教学管理及学院网站信息发布。积极开展大思政课实践教学探索，参与策划和组织学院师生赴杨浦滨江、红色场馆等“行走的思政课”与国门一线实践调研活动，协助推进学院党支部标准化规范化建设与特色党日活动开展。",
        "主页": "https://mks.shcc.edu.cn/2024/0311/c1591a59404/page.htm"
    }
]

def main():
    records = []
    for idx, item in enumerate(FACULTY_DATA, start=1):
        rec = {
            "序号": idx,
            "姓名": item["姓名"],
            "职称": item["职称"],
            "邮箱": item["邮箱"],
            "学校": item["学校"],
            "学院": item["学院"],
            "采集人": item["采集人"],
            "采集日期": item["采集日期"],
            "简介": item["简介"],
            "主页": item["主页"]
        }
        records.append(rec)

    school = "上海海关学院"
    college = "马克思主义学院"
    task_id = "7e60482a"
    date_str = "20260921"
    timestamp_str = datetime.datetime.now().strftime("%H%M%S")
    filename = f"{school}_{college}_{date_str}_{timestamp_str}.xlsx"
    output_dir = os.path.join(os.path.dirname(__file__), "outputs")
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)

    # 1. Create DataFrames
    df = pd.DataFrame(records)
    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    df = df[columns]

    df_with_email = df[df['邮箱'].str.strip() != ''].copy()
    df_no_email = df[df['邮箱'].str.strip() == ''].copy()

    # Reset index for individual sheets
    df_with_email['序号'] = range(1, len(df_with_email) + 1)
    df_no_email['序号'] = range(1, len(df_no_email) + 1)

    sheet1_name = f"有邮箱学者 (共 {len(df_with_email)} 位)"
    sheet2_name = f"无邮箱学者 (共 {len(df_no_email)} 位)"

    # 2. Write Dual Sheet Excel
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df_with_email.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_no_email.to_excel(writer, sheet_name=sheet2_name, index=False)

    # 3. Styling with openpyxl
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    data_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = data_font
                    if cell.column in [1, 2, 3, 5, 6, 7, 8]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col[:15]:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str.encode('gbk', 'ignore')))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    wb.save(filepath)
    print(f"🎉 Standard Dual-Sheet Excel saved to: {filepath}")
    print(f"  {sheet1_name}")
    print(f"  {sheet2_name}")
    print(f"  Total teachers: {len(df)} (Strictly Conserved: {len(df_with_email)} + {len(df_no_email)} = {len(df)})")

    # 4. Sync with Backend
    sync_url = f"http://127.0.0.1:8000/api/task/{task_id}/sync"
    sync_payload = {
        "records": records,
        "stats": {
            "total": len(records),
            "with_email": len(df_with_email),
            "no_email": len(df_no_email),
            "email_rate": round(len(df_with_email) / len(records) * 100, 1) if records else 0
        },
        "download_url": f"/download/{filename}",
        "progress": 100
    }

    req = urllib.request.Request(
        sync_url,
        data=json.dumps(sync_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )

    try:
        with urllib.request.urlopen(req) as resp:
            resp_data = resp.read().decode("utf-8")
            print(f"✅ Sync API Response: {resp_data}")
    except Exception as e:
        print(f"❌ Sync API Error: {e}")

if __name__ == "__main__":
    main()
