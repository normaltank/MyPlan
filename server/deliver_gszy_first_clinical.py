import os
import sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
import json
import datetime
import re
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import requests

TASK_ID = "33477c02e78a403d852186f3b44070cb"
SCHOOL = "甘肃中医药大学"
COLLEGE = "第一临床医学院"
AGENT_NAME = "Antigravity"
DATE_STR = datetime.datetime.now().strftime("%Y-%m-%d")
OUTPUT_DIR = "server/outputs"

def clean_cell(val):
    if val is None:
        return ""
    val_str = str(val)
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub('', val_str).strip()

def style_workbook(wb):
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    data_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = data_font
                    if cell.column in [1, 2, 3, 5, 6, 7, 8]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col[:20]:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str.encode('gbk', 'ignore')))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    raw_path = r"C:\Users\27451\.gemini\antigravity\brain\68366d88-1671-4d69-a94c-bb0e76794b0f\.system_generated\steps\209\output.txt"
    with open(raw_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line.startswith('{') and line.endswith('}'):
                raw_data = json.loads(line)
                break

    teachers = raw_data.get('teachers', [])
    deduped = {}
    for t in teachers:
        name = t['name']
        if name not in deduped:
            deduped[name] = t
        else:
            existing = deduped[name]
            if not existing.get('email') and t.get('email'):
                deduped[name] = t
            elif len(t.get('bio', '')) > len(existing.get('bio', '')):
                if existing.get('email') and not t.get('email'):
                    t['email'] = existing['email']
                deduped[name] = t

    unique_teachers = list(deduped.values())
    unique_teachers.sort(key=lambda x: x['name'])

    records = []
    for idx, t in enumerate(unique_teachers, start=1):
        records.append({
            "序号": idx,
            "姓名": clean_cell(t.get("name", "")),
            "职称": clean_cell(t.get("title", "教师")),
            "邮箱": clean_cell(t.get("email", "")),
            "学校": SCHOOL,
            "学院": COLLEGE,
            "采集人": AGENT_NAME,
            "采集日期": DATE_STR,
            "简介": clean_cell(t.get("bio", "")),
            "主页": clean_cell(t.get("url", ""))
        })

    df = pd.DataFrame(records)
    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    df_with_email = df[df['邮箱'].fillna('').astype(str).str.strip() != ''].copy().reset_index(drop=True)
    df_no_email = df[df['邮箱'].fillna('').astype(str).str.strip() == ''].copy().reset_index(drop=True)

    df_with_email['序号'] = range(1, len(df_with_email) + 1)
    df_no_email['序号'] = range(1, len(df_no_email) + 1)

    w_email = len(df_with_email)
    wo_email = len(df_no_email)
    total_college = len(df)
    assert w_email + wo_email == total_college, "双 Sheet 守恒校验失败！"

    # 1. Export college-specific Excel
    excel_college = os.path.join(OUTPUT_DIR, f"{SCHOOL}_{COLLEGE}.xlsx")
    with pd.ExcelWriter(excel_college, engine='openpyxl') as writer:
        df_with_email.to_excel(writer, sheet_name="教师列表", index=False)
        df_no_email.to_excel(writer, sheet_name="无邮箱教师列表", index=False)
    wb_col = openpyxl.load_workbook(excel_college)
    style_workbook(wb_col)
    wb_col.save(excel_college)
    print(f"✅ 学院专属表保存: {excel_college}")

    # 2. Merge into university-level Excel
    excel_univ = os.path.join(OUTPUT_DIR, f"{SCHOOL}.xlsx")
    if os.path.exists(excel_univ):
        try:
            old_with = pd.read_excel(excel_univ, sheet_name="教师列表")
            old_no = pd.read_excel(excel_univ, sheet_name="无邮箱教师列表")
            # Remove any existing records for this college
            old_with = old_with[old_with['学院'] != COLLEGE]
            old_no = old_no[old_no['学院'] != COLLEGE]
            
            merged_with = pd.concat([old_with, df_with_email], ignore_index=True)
            merged_no = pd.concat([old_no, df_no_email], ignore_index=True)
        except Exception:
            merged_with = df_with_email
            merged_no = df_no_email
    else:
        merged_with = df_with_email
        merged_no = df_no_email

    merged_with['序号'] = range(1, len(merged_with) + 1)
    merged_no['序号'] = range(1, len(merged_no) + 1)

    with pd.ExcelWriter(excel_univ, engine='openpyxl') as writer:
        merged_with.to_excel(writer, sheet_name="教师列表", index=False)
        merged_no.to_excel(writer, sheet_name="无邮箱教师列表", index=False)
    wb_univ = openpyxl.load_workbook(excel_univ)
    style_workbook(wb_univ)
    wb_univ.save(excel_univ)
    print(f"✅ 大学汇总表保存: {excel_univ} (全校总人数: {len(merged_with) + len(merged_no)}, 有邮箱: {len(merged_with)}, 无邮箱: {len(merged_no)})")

    # 3. Update root directory 教师列表.xlsx
    root_master = "教师列表.xlsx"
    if os.path.exists(root_master):
        try:
            master_df = pd.read_excel(root_master)
            # Remove existing for this school & college
            master_df = master_df[~((master_df['学校'] == SCHOOL) & (master_df['学院'] == COLLEGE))]
            updated_master = pd.concat([master_df, df], ignore_index=True)
            updated_master['序号'] = range(1, len(updated_master) + 1)
            updated_master.to_excel(root_master, index=False)
            print(f"✅ 根目录总库更新: {root_master} (总数: {len(updated_master)})")
        except Exception as e:
            print(f"⚠️ 根目录总库更新跳过: {e}")

    # 4. Sync to backend API
    sync_url = f"http://127.0.0.1:8000/api/task/{TASK_ID}/sync"
    sync_payload = {
        "records": records,
        "stats": {
            "total": total_college,
            "with_email": w_email,
            "no_email": wo_email,
            "email_rate": round(w_email / total_college * 100, 1),
            "school": SCHOOL,
            "college": COLLEGE
        },
        "download_url": f"/download/{SCHOOL}_{COLLEGE}.xlsx",
        "progress": 100,
        "worker_id": "Antigravity (IDE Agent)",
        "lease_token": "70fcdb0998b34afd8ed104ebf37bef54"
    }

    try:
        r = requests.post(sync_url, json=sync_payload, timeout=10)
        print(f"📡 后台同步响应 ({r.status_code}): {r.text[:200]}")
    except Exception as e:
        print(f"⚠️ 后台同步异常: {e}")

    print(json.dumps({
        "task_id": TASK_ID,
        "school": SCHOOL,
        "college": COLLEGE,
        "status": "completed",
        "total": total_college,
        "with_email": w_email,
        "excel": f"server/outputs/{SCHOOL}_{COLLEGE}.xlsx"
    }, ensure_ascii=False))

if __name__ == "__main__":
    main()
