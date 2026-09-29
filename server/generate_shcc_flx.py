import os
import sys
import io
import json
import urllib.request
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

RECORDS = [
    {
        "姓名": "祝少春",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "祝少春，男，中共党员，上海海关学院海关法律系副主任、副教授，全国反走私综合治理调查研究中心（NARC）特聘研究员，全国反走私综合治理调查研究中心上海分中心负责人，中华全国工商业联合会涉案企业合规第三方监督评估机制管理委员会第三方机制专业人员，上海市人民检察院第三分院跨行政区划案件特聘专家、听证员，上海市法学会海关法研究会副秘书长、理事，兼职律师。主要从事走私犯罪、反走私综合治理、海关缉私、海关监管领域教学、科研及干部培训。2020年，受海关总署缉私局委托，负责全国海关缉私系统刑事法制业务比武竞赛第一、三战区初赛和全国总决赛的命题及组织工作，获海关总署缉私局来函表扬。2017年9月，被联合国环境规划署和联合国亚洲及太平洋经济社会委员会授予“保护臭氧层突出贡献奖”。2022年，获评上海市课程思政教学名师。",
        "主页": "http://flx.shcc.edu.cn/2018/0314/c1902a14776/page.htm"
    },
    {
        "姓名": "朱秋沅",
        "职称": "教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "朱秋沅，法学博士，复旦大学法学博士后，上海海关学院海关法律系教授，硕士生导师，兼职律师；德国明斯特大学访问学者及其“海关与国际贸易法”硕士项目的“中国海关法律制度”课程的授课专家，德国哈尔茨大学公共管理专业常规课程“国际治理与区域化”的授课专家。同时，其为上海市法学会海关法研究会秘书长，中国欧洲学会欧洲法律研究会理事，上海市WTO法研究会理事，上海市法学会国际法研究会理事以及上海市法学会欧盟法研究会理事。\n\n其主要教学、科研领域为国际经济法、国际海关法、知识产权海关保护、原产地规则等，主持了《知识产权海关保护》署级精品课程（省部级）与市级重点课程建设，主持了《国际经济法》市级全英示范课程的建设工作。其所负责的上海海关学院法学专业教学体系改革获2017年上海市高等教育教学成果二等奖和校级教学成果一等奖。同时，其主持或深度参与了多项国家级或省部级项目，也深度参与了多项海关总署各部门与国家口岸管理办公室等的横向实务课题，并获得高度评价。其立有《国际海关法》《欧盟关税法及其壁垒性研究》《知识产权边境保护原理与实案》等专著5部、合著1部，编写海关与国际贸易法领域的教材8部；发表了论文50余篇，译作2篇，多篇发表于CSSCI系列刊物与北大版核心期刊，其中数篇被人大复印资料全文转载。其所主笔或深度参与的多篇论文曾被系统内《领导参阅》与《外事参阅件》所采用。\n\n2015年曾与德国AWA AUSSENWIRTSCHAFTS-AKADEMIE GmbH合作，在慕尼黑，明斯特等地对德国保时捷、罗氏等数十家著名企业提供关于“Chinese Customs Law and relevant Chinese foreign trade provisions”和“China-EU: Mutual recognition of Authorized Economic Operators”的实务培训。2016年曾在慕尼黑BMW全球总部的国际税务、海关与出口管制部进行了In-house的专题中国关务研讨和咨询。2018年曾赴深圳原产地管理办公室（海关总署原产地办的三个外脑机构之一）调研原产地实务工作。2015年间在上海欣海报关有限公司挂职，任副总经理。现为锦天城律师事务所高级顾问。\n\n其曾获“2017年度宝钢优秀教师奖”、“2012年度上海市育才奖”、2018年度“上海市教育系统三八红旗手”称号、校级“优秀教师”，第七届全国法律英语大赛“优秀指导教师奖”等。所指导的学生实践项目曾获第二届“知行杯”特等奖。",
        "主页": "http://flx.shcc.edu.cn/2016/0112/c1902a3982/page.htm"
    },
    {
        "姓名": "万曙春",
        "职称": "教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "万曙春，女，中共党员，上海海关学院海关法律系教授，法学博士，MPA硕士生导师。法国艾克斯马赛大学法学院访问学者、上海市青年法学法律人才库成员、上海市法学会宪法学学会和海关法学会理事。上海交通大学法学博士，上海社会科学院法学硕士；留学法国，法国巴黎第一大学-先贤祠-索邦大学欧洲法硕士。曾挂职上海外高桥保税区海关党组成员、副关长，法国国立海关（图尔昆）学校“现代海关管理”国际高级研修班结业。\n      研究领域：海关法与比较海关法、海关行政法治与缉私制度。发表专著《宪法视角下的海关缉私惩罚制度》，译著《海关法学-欧盟与法国海关法》（法语合译），编写教材《海关法精选案例实验教程》、《海关法文献集疏》；参编《海关法》、《比较海关法》和海关总署公务员初任培训教材编写及报关协会报关年度报告法规编纂等；主持国家社科重大子课题、上海哲社、上海市教委创新课题及其他课题10余项；在《人民日报（内参）》、《法学》、《政治与法律》、《欧洲与国际税法期刊》等国内外核心期刊及海关相关刊物发表决策咨询和论文30余篇。主持上海市市级全英示范课建《海关法》及多项院级教改项目。主讲海关法、海关法律综合实训、行政法与行政诉讼法等课程。",
        "主页": "http://flx.shcc.edu.cn/2016/0111/c1902a3948/page.htm"
    },
    {
        "姓名": "娄万锁",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "娄万锁，浙江台州人，上海海关学院海关法律系副教授，北京德和衡（上海）律师事务所律师。1999年7月毕业于复旦大学法学院法制史专业，获法学硕士学位；2012年6月毕业于复旦大学国际关系与公共事务学院政治学理论专业，获法学博士学位。2015年3月至9月英国诺丁汉大学访问学者。\n\n曾担任浙江省三门县小雄中学教师，上海海关学院（海关）法律系副主任、教工党支部书记等职。\n\n主讲课程：法理学、政治学、海关法、关贸律师实务（本科）\n\n             政治学、依法治国与中国特色社会主义政治建设（研究生）\n\n研究方向：海关发展战略，海关廉政、海关法，公司法\n\n主要学术成果：\n\n1、专著《中国海关改革的政治学分析》，上海人民出版社2015年5月；\n\n2、参编《欧盟经济行政法通论》，东方出版中心2000年11月；\n\n3、参编《案例海关法》，立信会计出版社2007年4月；\n\n4、参编《海关法理论与实践》，立信会计出版社2008年3月；\n\n5、参编《海关法学》，中国海关出版社2010年4月；\n\n6、《试论一人公司法律制度》，《社会科学战线》1999年3月；\n\n7、《论加工贸易中的海关行政许可》，《海关贯彻实施行政许可法理论与实践》2004年（该文获中国海关学会、海关总署政法司举办的征文二等奖）；\n\n8、《全球化与海关职能转变》，《上海海关学院学报》2008年6月；\n\n9、《贸易便利化理论根源探析》，《上海海关学院学报》2010年3月；\n\n10、《国际制度视角下区域贸易争端解决机制研究》，《兰州学刊》2012年3月；\n\n11、《中国海关廉政建设中的政治象征研究》，《云南行政学院学报》2012年3月；\n\n12、《论中国海关的产生》，《上海海关学院学报》2012年10月；\n\n13、《论新时期海关发展战略的几个基本问题》，《海关与经贸研究》2014年3月；\n\n14、《转向权利本位：“负面清单”管理的法理分析》，《海关法评论》2014年9月；\n\n15、《“长江经济带”通关一体化中的海关法律问题探析》，《海关法评论》2016年12月；\n\n16、《中国海关治理七十年：历程、经验及展望》，《海关与经贸研究》2019年11月；\n\n17、参与完成国家社科基金重点项目“加强我国海关管理的目标与机制研究（13AGL006）”；\n\n18、主持完成上海市教委科研创新重点项目“垂直领导部门与地方政府关系研究—以海关为例”（14ZS165）；\n\n19、参与完成海关总署项目“新时期海关发展战略研究”（2014）\n\n20、参与完成海关总署项目“新时代海关‘两步走’战略研究”（2019）；\n\n21、主持完成上海海关学院项目“上海自贸区政府治理体制机制研究”（2014）；\n\n22、主持完成上海海关学院项目“‘长江经济带’大通关建设协同治理研究”（2017）。",
        "主页": "http://flx.shcc.edu.cn/2016/0108/c1902a3922/page.htm"
    },
    {
        "姓名": "刘海燕",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "刘海燕，女，上海海关学院海关法律系副教授。毕业于大连海事大学国际经济法专业（海商法方向），获硕士学位，现于上海海事大学在读博士。主要从事国际经济法、海关国际法、海商法的教学与研究。主讲《国际私法》、《经济法》、《海商法》等多门课程。曾先后在《中国航海》、《财经问题研究》、《辽宁大学学报》、《理论界》、《水运管理》等核心刊物上发表多篇理论研究文章，并参与编写8本教材和学术著作，参与近十项课题的研究。",
        "主页": "http://flx.shcc.edu.cn/2016/0111/c1902a3974/page.htm"
    },
    {
        "姓名": "吴展",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "吴展，汉族，安徽萧县人，上海海关学院海关法律系副教授，法学博士、海关法教研室主任、海关与公共管理学院硕士研究生导师、兼职律师。主持国家社会科学基金重大招投标项目子课题、教育部人文社会科学青年基金项目等科研项目11项、参与国家社会科学基金重大招投标项目、上海市教委、海关总署署级课题等科研项目18项。在《比较法研究》、《政治与法律》、《国际商务研究》等专业期刊发表论文30余篇，其中多篇论文被中国人民大学复印资料中心教授全文转载。参与编撰《海关法学》、《海关行政法原理与实务》、《海关行政法实训案例教程》、《案例海关法教程》、《司法法学》等论著多部，独立撰写《海关行政处罚：理论、实务与合规指引》、《海关法律英语教程》两部。",
        "主页": "http://flx.shcc.edu.cn/2016/0108/c1902a3923/page.htm"
    },
    {
        "姓名": "王珉",
        "职称": "副教授",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "王珉：上海海关学院海关法律系副教授，教研室主任，硕士研究生导师，复旦大学法学博士，上海市“晨光学者”，上海市法学会海关法研究会副秘书长。主要研究方向为国际经济法、比较海关法、数字贸易法。\n\n主持上海高等学校一流本科课程《国际海关法》的建设工作，主研或主持国家级、省部级项目以及海关总署各部门研究课题和横向课题等20余项。出版专著1部《全球治理视域下的国际海关法研究》、合著3部，在World Customs Journal（世界海关组织（WCO）官方期刊《世界海关学刊》）等英文期刊和《国际贸易》《东北亚论坛》《太平洋学报》《河北法学》等CSSCI期刊（含扩展版）及其他学术期刊发表《DEPA规则对数字贸易海关监管的启示及中国对策》《中国贸易便利化发展战略——基于RCEP成员之间区域自贸协定的比较分析》《丝绸之路经济带海关国际合作法律机制的构建》《国际海关便利通关制度最佳实践经验、发展趋势与中国因应之策》等学术论文20余篇。曾受邀参与世界海关组织（WCO）第十次海关学术研究与发展伙伴关系会议（PICARD）并作主题演讲。",
        "主页": "http://flx.shcc.edu.cn/2016/0108/c1902a3924/page.htm"
    },
    {
        "姓名": "范筱静",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "范筱静，女，1979年10月出生，山东烟台人，法学博士研究生，上海海关学院海关法律系讲师，兼职律师。\n\n 2002年7月毕业于山东烟台大学法学专业，2006年3月毕业于上海对外贸易大学国际经济法专业，获法学硕士学位；2013年6月毕业于上海交通大学诉讼法专业，获法学博士学位。自2006年7月开始在上海海关学院工作，先后担任过法律系辅导员、法律系主任助理、法律系党总支副书记、科研处主任助理等职务，现担任法律系专任教师，负责民事诉讼法、法理学、海关法、国际商事仲裁、电子商务法等课程的教学和科研，主要研究方向为纠纷解决机制、(跨境)电子商务法律问题、海关法等。参编《海关国际法》、《海关国际商务法》、《海关法学》、《海关法理论与实践》、《一带一路战略与海关国际合作法律机制》等著作，参与或主持国家哲社、海关总署等各级各类课题十余项，并在各级各类学术期刊上发表论文十几篇。2005年4月取得律师执业证书，现任北京德和衡（上海）律师事务所兼职律师。",
        "主页": "http://flx.shcc.edu.cn/2016/0111/c1902a3968/page.htm"
    },
    {
        "姓名": "朱江华",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "朱江华，1974年6月出生，男，汉族，江西泰和人，中共党员，法学硕士学位，上海海关学院海关法律系讲师，实验室主任，主要从事诉讼法、海关缉私等领域的教学研究工作，在公开刊物发表论文多篇，并参与《海关缉私总论》、《案例海关法教程》、《海关法释解》等教材或著作的编写。",
        "主页": "http://flx.shcc.edu.cn/2018/0314/c1902a14778/page.htm"
    },
    {
        "姓名": "王永亮",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "王永亮，男，上海海关学院海关法律系讲师。1999年毕业于郑州大学法学院，获法学学士学位；2004年毕业于郑州大学法学院，获民商法专业硕士学位；2014年6月毕业于华东政法大学，获法学博士学位。2000年至2001年在洛阳市开物律师事务所任律师助理。2004年7月进入上海市金山区人民法院工作，2007年在上海海事法院轮岗锻炼半年，2014年7月起于上海海关学院法律系工作。主要研究领域为民法、商法等。2015年参与海关总署署级课题撰写，并担任上海海关学院法律顾问。",
        "主页": "http://flx.shcc.edu.cn/2016/0108/c1902a3927/page.htm"
    },
    {
        "姓名": "陈阳",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "陈阳，女，上海海关学院海关法律系讲师，民商法教学团队负责人。博士毕业于华东政法大学，2014年入职上海海关学院。承担《民法》《合同法》《国际商法》《中国法律史》等课程的教学任务，长期从事民商法、国际商法、国际海关法领域的研究工作。在《山东社会科学》《理论月刊》等期刊发表论文数篇，出版1部婚姻法领域专著，参编多部著作，主持并参与多项课题。",
        "主页": "http://flx.shcc.edu.cn/2020/0316/c1902a42671/page.htm"
    },
    {
        "姓名": "纪康",
        "职称": "讲师",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "纪康，上海海关学院海关法律系讲师，法学博士，主要研究方向：刑法学、走私犯罪，在《首都师范大学学报（社会科学版）》《刑法论丛》《法律方法》等cssci来源期刊上发表论文数篇，主持、参与国家级、省部级课题多项。",
        "主页": "http://flx.shcc.edu.cn/2024/0619/c1902a60972/page.htm"
    },
    {
        "姓名": "李航",
        "职称": "助教",
        "邮箱": "",
        "学校": "上海海关学院",
        "学院": "法学教学部",
        "采集人": "Antigravity Subagent (Browser)",
        "采集日期": "2026-09-21",
        "简介": "李航，上海海关学院海关法律系助教，华东政法大学经济法学博士，博士期间研究方向为金融法，公开发表核心论文数篇，参与修订、编纂教材多部，参与省部级课题、期货交易所横向课题等多项。",
        "主页": "http://flx.shcc.edu.cn/2025/1009/c1902a66972/page.htm"
    }
]

def main():
    school = "上海海关学院"
    college = "法学教学部"
    task_id = "61b62da3"
    
    # 1. Output directory
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    output_dir = os.path.join(base_dir, "server", "outputs")
    os.makedirs(output_dir, exist_ok=True)
    
    timestamp = pd.Timestamp.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{school}_{college}_{timestamp}.xlsx"
    filepath = os.path.join(output_dir, filename)
    
    # 2. DataFrame and Split
    df = pd.DataFrame(RECORDS)
    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    
    df_with_email = df[df['邮箱'].fillna('').astype(str).str.strip() != ""].copy().reset_index(drop=True)
    df_no_email = df[df['邮箱'].fillna('').astype(str).str.strip() == ""].copy().reset_index(drop=True)
    
    df_with_email['序号'] = range(1, len(df_with_email) + 1)
    df_no_email['序号'] = range(1, len(df_no_email) + 1)
    
    for c in columns:
        if c not in df_with_email.columns:
            df_with_email[c] = ""
        if c not in df_no_email.columns:
            df_no_email[c] = ""
            
    df_with_email = df_with_email[columns]
    df_no_email = df_no_email[columns]
    
    sheet1_name = f"有邮箱学者 (共 {len(df_with_email)} 位)"
    sheet2_name = f"无邮箱学者 (共 {len(df_no_email)} 位)"
    
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df_with_email.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_no_email.to_excel(writer, sheet_name=sheet2_name, index=False)
        
    # 3. Styling with openpyxl
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    data_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = data_font
                    if cell.column in [1, 2, 3, 5, 6, 7, 8]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col[:15]:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str.encode('gbk', 'ignore')))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    wb.save(filepath)
    print(f"🎉 Standard Dual-Sheet Excel saved to: {filepath}")
    print(f"  {sheet1_name}")
    print(f"  {sheet2_name}")
    print(f"  Total teachers: {len(df)} (Strictly Conserved: {len(df_with_email)} + {len(df_no_email)} = {len(df)})")
    
    # 4. Sync with Backend
    sync_url = f"http://127.0.0.1:8000/api/task/{task_id}/sync"
    sync_payload = {
        "records": RECORDS,
        "stats": {
            "total": len(RECORDS),
            "with_email": len(df_with_email),
            "no_email": len(df_no_email),
            "email_rate": round(len(df_with_email) / len(RECORDS) * 100, 1) if RECORDS else 0
        },
        "download_url": f"/download/{filename}",
        "progress": 100
    }
    
    req = urllib.request.Request(
        sync_url,
        data=json.dumps(sync_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    
    try:
        with urllib.request.urlopen(req) as resp:
            resp_data = resp.read().decode("utf-8")
            print(f"✅ Sync API Response: {resp_data}")
    except Exception as e:
        print(f"❌ Sync API Error: {e}")

if __name__ == "__main__":
    main()
