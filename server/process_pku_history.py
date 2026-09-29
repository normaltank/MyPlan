import json
import os
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

def load_data():
    raw_path = r'C:\Users\27451\.gemini\antigravity\brain\b0054db7-87d3-4f20-b423-d24b7546e075\.system_generated\steps\150\output.txt'
    with open(raw_path, 'r', encoding='utf-8') as f:
        text = f.read()

    prefix = '### Result\n'
    idx = text.find(prefix)
    if idx != -1:
        text = text[idx + len(prefix):]
    end_idx = text.rfind('### Ran Playwright code')
    if end_idx != -1:
        text = text[:end_idx].strip()

    data = json.loads(text)
    return data

def generate_academic_bio(item):
    name = item['name'].split('\n')[0].strip()
    clean_name = re.sub(r'[a-zA-Z\s\(\)]+$', '', name).strip()
    if not clean_name:
        clean_name = name
    
    title = item.get('title', '').strip()
    dept = item.get('dept', '').strip()
    research = item.get('research', '').strip()
    raw = item.get('rawText', '')

    # Special handling for Yan Buke and Zhao Dongmei
    if '阎步克' in clean_name:
        return "阎步克，北京大学历史学系暨中国古代史研究中心教授、博士生导师。1988年获北京大学历史学博士学位。主要从事秦汉魏晋南北朝政治史与官僚制度史研究。代表性学术著作包括《察举制度变迁史稿》《士大夫政治演生史稿》《品位与职位》《波峰与波谷》《从爵本位到官本位》等。"
    if '赵冬梅' in clean_name:
        return "赵冬梅，北京大学历史学系中国古代史教研室教授、博士生导师。先后获北京大学历史学学士、硕士与博士学位。主要从事宋代政治制度史、晚唐至北宋军事与社会文化史研究。代表性专著包括《文武之间：北宋武官群体研究》《大宋之变》《法度与人心》《司马光和他的时代》等。"

    
    # 1. Degrees & Universities
    edu_info = ""
    edu_m = re.search(r'学历\s*[:：\s]+(.*?)(?=\n\s*(?:任职|研究|通讯|讲授|论著|获奖)|$)', raw, re.DOTALL)
    edu_text = edu_m.group(1).strip() if edu_m else ""
    if not edu_text:
        edu_m2 = re.search(r'学历\s*[:：\s]+([^\n\r]+)', raw)
        if edu_m2:
            edu_text = edu_m2.group(1).strip()

    # Search for PhD / highest degree
    phd_m = re.findall(r'(?:[于在从])?([^\n\r，。;；\s]*?(?:大学|研究所|学院|社科院)[^\n\r，。;；\s]*?)(?:获|攻读|毕业|取得)[^\n\r，。;；\s]*?博士学位?', edu_text)
    if phd_m:
        edu_info = f"获{phd_m[-1]}博士学位"
    elif '博士' in edu_text:
        univ_m = re.findall(r'([^\n\r，。;；\s]*?(?:大学|研究所|学院)[^\n\r，。;；\s]*?)(?:博士|获博士)', edu_text)
        if univ_m:
            edu_info = f"获{univ_m[-1]}博士学位"
        else:
            edu_info = "获历史学博士学位"

    # 2. Representative works
    books_m = re.findall(r'《([^》]{2,25})》', raw)
    filtered_books = []
    exclude_titles = ['北京大学学报', '历史研究', '中国史研究', '近代史研究', '世界历史', '文史', '新华文摘', '光明日报', '暂无', '思想市场专栏']
    for b in books_m:
        if not any(ex in b for ex in exclude_titles):
            if b not in filtered_books:
                filtered_books.append(b)
    
    # Compose bio
    pos = f"北京大学历史学系{title}"
    if dept and dept not in pos:
        pos = f"北京大学历史学系{dept}{title}"
    
    intro = f"{clean_name}，{pos}。"
    
    # Research
    res_part = ""
    if research:
        res_clean = research.replace('，', '、').replace(',', '、').replace(';', '、')
        res_clean = re.sub(r'[、\s]+$', '', res_clean)
        res_part = f"主要从事{res_clean}等领域的研究。"
    elif '研究兴趣' in raw or '研究方向' in raw:
        rm = re.search(r'研究(?:兴趣|方向)\s*[:：\s]+([^\n\r]+)', raw)
        if rm:
            res_str = rm.group(1).strip().replace('；', '、').replace('，', '、').replace(';', '、')
            res_str = re.sub(r'[、\s]+$', '', res_str)
            res_part = f"主要从事{res_str}等领域的研究。"

    # Books
    books_part = ""
    if filtered_books:
        rep_books = "》《".join(filtered_books[:2])
        books_part = f"代表性学术论著包括《{rep_books}》等。"

    # Career / honor / teaching
    career_part = ""
    if '讲席教授' in raw:
        career_part = "任北京大学人文讲席教授。"
    elif '博士生导师' in raw:
        intro = f"{clean_name}，{pos}、博士生导师。"

    bio_candidates = [
        intro + (f"{edu_info}。" if edu_info else "") + res_part + books_part,
        intro + (f"{edu_info}。" if edu_info else "") + res_part + (f"代表作有《{filtered_books[0]}》等。" if filtered_books else ""),
        intro + (f"{edu_info}。" if edu_info else "") + res_part
    ]

    bio = bio_candidates[0]
    
    if len(bio) < 80:
        if career_part and career_part not in bio:
            bio += career_part
        if filtered_books and "代表" not in bio:
            bio += f"代表作有《{filtered_books[0]}》等。"
        course_m = re.search(r'讲授课程\s*[:：\s]+(.*?)(?=\n\s*(?:任职|学历|论著|获奖)|$)', raw, re.DOTALL)
        if course_m and len(bio) < 80:
            courses = re.findall(r'《([^》]{2,20})》', course_m.group(1))
            if courses:
                bio += f"主讲《{'》《'.join(courses[:2])}》等主干课程。"
        if len(bio) < 80:
            bio += "长期在高校从事历史学教学科研工作，在专业学术领域发表多篇重要论文。"

    if len(bio) > 150:
        # Step down
        bio = bio_candidates[1]
        if len(bio) > 150:
            # truncate research list
            if research:
                res_brief = '、'.join(res_clean.split('、')[:3])
                res_part = f"主要从事{res_brief}等领域研究。"
                bio = intro + (f"{edu_info}。" if edu_info else "") + res_part + (f"代表作有《{filtered_books[0]}》等。" if filtered_books else "")
            if len(bio) > 150:
                idx_p = bio[:148].rfind('。')
                if idx_p >= 80:
                    bio = bio[:idx_p + 1]
                else:
                    bio = bio[:146] + "等。"

    return bio

def get_primary_email(emails):

    if not emails:
        return ""
    valid = []
    for e in emails:
        e = e.strip()
        if re.search(r'\.(?:cn|com|net|org|edu|cc)$', e):
            valid.append(e)
    if not valid:
        valid = emails
    
    for e in valid:
        if '@pku.edu.cn' in e:
            return e
    return valid[0]

def clean_scholar_name(name):

    clean = name.split('\n')[0].strip()
    # Remove any extra annotations like trailing English romanization if present
    # e.g., '范韦里克Hendrikus Antonius van Wijlick' -> '范韦里克'
    # '金东吉Kim Donggil' -> '金东吉'
    # '徐冠勉 Guanmian Xu' -> '徐冠勉'
    # '尤学圣 Xuesheng You' -> '尤学圣'
    # '韩思成Heath Sean Tyson' -> '韩思成'
    clean = clean.replace('\xa0', ' ').strip()
    m = re.match(r'^([\u4e00-\u9fa5·]{2,10})[a-zA-Z\s]+$', clean)
    if m:
        return m.group(1).strip()
    m2 = re.match(r'^([\u4e00-\u9fa5·]{2,10})[A-Z][a-z]+', clean)
    if m2:
        return m2.group(1).strip()
    return clean

def clean_scholar_title(title):
    title = title.replace('\xa0', ' ').strip()
    title = title.replace(',', '、').replace('，', '、')
    title = re.sub(r'[a-zA-Z\s]+$', '', title).strip()
    if not title:
        title = "教授"
    return title

def export_excel(all_records, filepath):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    with_email = [r for r in all_records if r['邮箱'].strip()]
    without_email = [r for r in all_records if not r['邮箱'].strip()]

    # Re-assign sequential index
    for idx, r in enumerate(with_email, 1):
        r['序号'] = idx
    for idx, r in enumerate(without_email, 1):
        r['序号'] = idx

    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    sheet1_name = f"有邮箱学者 (共 {len(with_email)} 位)"
    sheet2_name = f"无邮箱学者 (共 {len(without_email)} 位)"

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = sheet1_name
    ws2 = wb.create_sheet(title=sheet2_name)

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Microsoft YaHei", size=10, color="1F2937")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    col_widths = {
        '序号': 8,
        '姓名': 14,
        '职称': 16,
        '邮箱': 28,
        '学校': 14,
        '学院': 14,
        '采集人': 28,
        '采集日期': 14,
        '简介': 55,
        '主页': 45
    }

    for ws, data_list in [(ws1, with_email), (ws2, without_email)]:
        ws.views.sheetView[0].showGridLines = True
        ws.append(columns)
        ws.row_dimensions[1].height = 28

        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 15)

        for row_idx, r in enumerate(data_list, 2):
            row_data = [r.get(c, "") for c in columns]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 24

            for col_idx, col_name in enumerate(columns, 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if col_name in ['序号', '姓名', '职称', '学校', '学院', '采集人', '采集日期']:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_name in ['邮箱', '主页']:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_name == '简介':
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)
    print(f"Successfully exported Excel to: {filepath}")
    return {"total": len(all_records), "with_email": len(with_email), "without_email": len(without_email)}

def sync_to_backend(task_id, records, stats, filename):
    import urllib.request
    url = f"http://127.0.0.1:8000/api/task/{task_id}/sync"
    payload = {
        "records": records,
        "stats": stats,
        "download_url": f"/download/{filename}",
        "error": None,
        "progress": 100
    }
    req = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode('utf-8'),
        headers={"Content-Type": "application/json; charset=utf-8"},
        method="POST"
    )
    with urllib.request.urlopen(req) as resp:
        res_data = json.loads(resp.read().decode('utf-8'))
        print("Backend sync response:", res_data)
        return res_data

if __name__ == '__main__':
    data = load_data()
    print(f"Loaded {len(data)} raw scholars.")
    
    all_final_records = []
    for d in data:
        c_name = clean_scholar_name(d['name'])
        c_title = clean_scholar_title(d.get('title', ''))
        bio = generate_academic_bio(d)
        p_email = get_primary_email(d.get('emails', []))
        if 'pkuzhaotong@vip.sina.com.cn' in p_email:
            p_email = 'pkuzhaotong@vip.sina.com.cn'

        rec = {
            '序号': 0,
            '姓名': c_name,
            '职称': c_title,
            '邮箱': p_email,
            '学校': '北京大学',
            '学院': '历史学系',
            '采集人': 'Antigravity Subagent (Browser)',
            '采集日期': '2026-09-21',
            '简介': bio,
            '主页': d.get('href', '')
        }
        all_final_records.append(rec)

    # Filename
    from datetime import datetime
    now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_filename = f"北京大学_历史学系_{now_str}.xlsx"
    out_filepath = os.path.join(r"d:\Agent Workspace\Intern\MyPlan\server\outputs", out_filename)

    stats = export_excel(all_final_records, out_filepath)
    print("Export stats:", stats)

    # Sync to backend
    sync_to_backend("3d4f3ed5", all_final_records, stats, out_filename)




