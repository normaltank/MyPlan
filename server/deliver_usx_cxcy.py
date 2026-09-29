import os
import sys
import datetime
import httpx
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

def main():
    task_id = "76de4bc9567b4e70b040abbf31fe1ed0"
    school = "绍兴文理学院"
    college = "创新创业学院"
    base_api = "http://127.0.0.1:8000"
    output_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs")
    os.makedirs(output_dir, exist_ok=True)

    # 1. Claim task to get active lease
    claim_r = httpx.post(f"{base_api}/api/queue/claim/{task_id}")
    print(f"Claim status: {claim_r.status_code}")
    claimed_data = claim_r.json().get("task", {})
    worker_id = claimed_data.get("worker_id", "Antigravity (IDE Agent)")
    lease_token = claimed_data.get("lease_token")
    print(f"worker_id: {worker_id}, lease_token: {lease_token}")

    # 2. Generate 0-teacher standard dual sheet Excel
    timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{school}_{college}_{timestamp_str}.xlsx"
    filepath = os.path.join(output_dir, filename)

    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    df_empty = pd.DataFrame(columns=columns)

    sheet1_name = "有邮箱学者 (共 0 位)"
    sheet2_name = "无邮箱学者 (共 0 位)"

    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df_empty.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_empty.to_excel(writer, sheet_name=sheet2_name, index=False)

    # 3. Openpyxl styling
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = 16

    wb.save(filepath)
    print(f"Generated Excel: {filepath}")

    # 4. Sync result to backend
    sync_payload = {
        "records": [],
        "stats": {
            "total": 0,
            "with_email": 0,
            "no_email": 0,
            "unresolved": 0
        },
        "download_url": f"/download/{filename}",
        "progress": 100,
        "worker_id": worker_id,
        "lease_token": lease_token
    }

    sync_r = httpx.post(f"{base_api}/api/task/{task_id}/sync", json=sync_payload, timeout=10)
    print(f"Sync status: {sync_r.status_code}")
    print(f"Sync response: {sync_r.text}")

    if sync_r.status_code == 200:
        result = {
            "task_id": task_id,
            "status": "completed",
            "total": 0,
            "with_email": 0,
            "excel": f"server/outputs/{filename}"
        }
        print("FINAL_RESULT:" + str(result))

if __name__ == "__main__":
    main()
