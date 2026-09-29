import os
import sys
import json
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

OUTPUT_DIR = "server/outputs"
with open("server/agent_queue.json", "r", encoding="utf-8") as f:
    data = json.load(f)

batch_id = "4f0dee1a"
tids = data["batch_tasks"][batch_id]

all_records = []
summary = []

print(f"=== PROCESSING BATCH {batch_id} (18 TASKS) ===")
for idx, tid in enumerate(tids, 1):
    t = data["tasks"][tid]
    school = t["school"]
    college = t["college"]
    down_url = t.get("download_url", "")
    records = t.get("records", [])
    
    # If records not stored directly in task, read from excel
    if not records and down_url:
        filename = down_url.replace("/download/", "")
        filepath = os.path.join(OUTPUT_DIR, filename)
        if os.path.exists(filepath):
            try:
                df1 = pd.read_excel(filepath, sheet_name=0)
                df2 = pd.read_excel(filepath, sheet_name=1)
                records = pd.concat([df1, df2], ignore_index=True).fillna('').to_dict(orient='records')
            except Exception as e:
                print(f"Error reading {filepath}: {e}")
                
    w_email = sum(1 for r in records if str(r.get("邮箱", "")).strip() != "")
    no_email = len(records) - w_email
    
    print(f"{idx:2d}. [{school}] {college}: 共 {len(records)} 人 (有邮箱 {w_email}, 无邮箱 {no_email})")
    summary.append({
        "school": school,
        "college": college,
        "total": len(records),
        "with_email": w_email,
        "no_email": no_email
    })
    all_records.extend(records)

print(f"\nTotal raw collected across 18 units: {len(all_records)}")

# Deduplication within batch (e.g. if same scholar appears multiple times with same name & email)
df_all = pd.DataFrame(all_records)
if not df_all.empty:
    # Normalize columns
    expected_cols = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    for c in expected_cols:
        if c not in df_all.columns:
            if c == '学校' and '高校' in df_all.columns:
                df_all['学校'] = df_all['高校']
            elif c == '主页' and '个人主页' in df_all.columns:
                df_all['主页'] = df_all['个人主页']
            elif c == '简介' and '个人简介' in df_all.columns:
                df_all['简介'] = df_all['个人简介']
            else:
                df_all[c] = ""
    
    # Ensure deduplication: for duplicate (学校, 学院, 姓名, 邮箱), keep first
    df_all = df_all.drop_duplicates(subset=['学校', '学院', '姓名', '邮箱'], keep='first').reset_index(drop=True)
    
    # Split into Sheet 1 & Sheet 2
    has_email_mask = df_all['邮箱'].fillna('').astype(str).str.strip() != ''
    df_s1 = df_all[has_email_mask].copy().reset_index(drop=True)
    df_s2 = df_all[~has_email_mask].copy().reset_index(drop=True)
    
    df_s1['序号'] = range(1, len(df_s1) + 1)
    df_s2['序号'] = range(1, len(df_s2) + 1)
    
    df_s1 = df_s1[expected_cols]
    df_s2 = df_s2[expected_cols]
    
    out_name = f"批量采集教师列表_合集_南京林业大学_江苏大学_18学院_20260922.xlsx"
    out_path = os.path.join(OUTPUT_DIR, out_name)
    
    sheet1_name = f"有邮箱学者 (共 {len(df_s1)} 位)"
    sheet2_name = f"无邮箱学者 (共 {len(df_s2)} 位)"
    
    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        df_s1.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_s2.to_excel(writer, sheet_name=sheet2_name, index=False)
        
    # Openpyxl styling
    wb = openpyxl.load_workbook(out_path)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    data_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = data_font
                    if cell.column in [1, 2, 3, 5, 6, 7, 8]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')
        
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col[:15]:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str.encode('gbk', 'ignore')))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)
            
    wb.save(out_path)
    print(f"\n🎉 Successfully created Master Merged Excel: {out_name}")
    print(f"Sheet 1: {sheet1_name}")
    print(f"Sheet 2: {sheet2_name}")
    print(f"Conserved Total: {len(df_s1) + len(df_s2)}")
    
    # Update agent_queue._batch_tasks
    data["batch_tasks"][f"{batch_id}_merged_file"] = out_name
    with open("server/agent_queue.json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print("Updated agent_queue.json batch merged link.")
