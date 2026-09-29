import os
import sys
import io
import json
import re
import datetime
import urllib.request
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

TEACHER_TITLES = {
    '孙浩': '教授',
    '黄丙志': '教授',
    '周欣': '副教授',
    '陈振海': '讲师',
    '张慧': '副教授',
    '刘宗沅': '讲师',
    '周国荣': '讲师',
    '谢晶': '副教授',
    '温朝柱': '副教授',
    '宗慧民': '副教授',
    '侯彩虹': '副教授',
    '姜晓依': '副教授',
    '王菲易': '副教授',
    '丁敏': '副教授',
    '林弘': '副教授',
    '王诗勇': '讲师',
    '路江林': '讲师',
    '宋杨': '讲师',
    '王静': '讲师',
    '于潇': '讲师',
    '严晓': '讲师',
    '王瑜': '讲师',
    '孔晓丹': '讲师',
    '王华强': '讲师',
    '金宏彬': '讲师',
    '刘文佳': '讲师',
    '俞钦文': '讲师',
    '周芮': '讲师',
    '郁姣娇': '讲师',
    '邓敏': '讲师',
    '姜玉瑶': '讲师',
    '李珍': '讲师',
    '郑夏': '讲师',
    '孔哲礼': '讲师',
    '魏彬彬': '讲师',
    '陈蓉': '讲师',
    '齐昕': '讲师',
}

# Order teachers logically: leadership -> professors -> associate professors -> lecturers
ORDER = [
    '孙浩', '黄丙志',
    '周欣', '张慧', '谢晶', '温朝柱', '宗慧民', '侯彩虹', '姜晓依', '王菲易', '丁敏', '林弘',
    '陈振海', '刘宗沅', '周国荣', '王诗勇', '路江林', '宋杨', '王静', '于潇',
    '严晓', '王瑜', '孔晓丹', '王华强', '金宏彬', '刘文佳', '俞钦文', '周芮',
    '郁姣娇', '邓敏', '姜玉瑶', '李珍', '郑夏', '孔哲礼', '魏彬彬', '陈蓉', '齐昕'
]

def main():
    json_path = r'C:\Users\27451\.gemini\antigravity\brain\cbf60285-1dc7-4899-a08d-bc1932c56318\scratch\teachers.json'
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    records = []
    idx = 1
    for name in ORDER:
        if name not in data:
            print(f"Warning: {name} not found in scraped data!")
            continue
        item = data[name]
        raw_text = item['text']
        # Clean CMS metadata headers
        clean_text = re.sub(r'^[^\n]+\n+发布者[：:][^\n]*\n+发布时间[：:][^\n]*\n+浏览次数[：:]\n*\d*\s*', '', raw_text).strip()
        # Clean leading noise if present
        clean_text = re.sub(r'^[\s\xa0]+', '', clean_text)
        
        record = {
            "序号": idx,
            "姓名": name,
            "职称": TEACHER_TITLES.get(name, item.get('title', '讲师')),
            "邮箱": item.get('email', '').strip(),
            "学校": "上海海关学院",
            "学院": "海关与公共管理学院",
            "采集人": "Antigravity Subagent (Browser)",
            "采集日期": "2026-09-21",
            "简介": clean_text,
            "主页": item['url']
        }
        records.append(record)
        idx += 1

    school = "上海海关学院"
    college = "海关与公共管理学院"
    task_id = "17291eaa"
    date_str = "20260921"
    timestamp_str = datetime.datetime.now().strftime("%H%M%S")
    filename = f"{school}_{college}_{date_str}_{timestamp_str}.xlsx"
    output_dir = os.path.join(os.path.dirname(__file__), "outputs")
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)

    # 1. Create DataFrames
    df = pd.DataFrame(records)
    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    df = df[columns]

    df_with_email = df[df['邮箱'].str.strip() != ''].copy()
    df_no_email = df[df['邮箱'].str.strip() == ''].copy()

    # Reset index for individual sheets
    df_with_email['序号'] = range(1, len(df_with_email) + 1)
    df_no_email['序号'] = range(1, len(df_no_email) + 1)

    sheet1_name = f"有邮箱学者 (共 {len(df_with_email)} 位)"
    sheet2_name = f"无邮箱学者 (共 {len(df_no_email)} 位)"

    # 2. Write Dual Sheet Excel
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df_with_email.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_no_email.to_excel(writer, sheet_name=sheet2_name, index=False)

    # 3. Styling with openpyxl
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    data_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            for cell in row:
                cell.border = thin_border
                if row_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = data_font
                    if cell.column in [1, 2, 3, 5, 6, 7, 8]:
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='left', vertical='center')

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col[:15]:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str.encode('gbk', 'ignore')))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    wb.save(filepath)
    print(f"🎉 Standard Dual-Sheet Excel saved to: {filepath}")
    print(f"  {sheet1_name}")
    print(f"  {sheet2_name}")
    print(f"  Total teachers: {len(df)} (Strictly Conserved: {len(df_with_email)} + {len(df_no_email)} = {len(df)})")

    # 4. Sync with Backend
    sync_url = f"http://127.0.0.1:8000/api/task/{task_id}/sync"
    sync_payload = {
        "records": records,
        "stats": {
            "total": len(records),
            "with_email": len(df_with_email),
            "no_email": len(df_no_email),
            "email_rate": round(len(df_with_email) / len(records) * 100, 1) if records else 0
        },
        "download_url": f"/download/{filename}",
        "progress": 100
    }

    req = urllib.request.Request(
        sync_url,
        data=json.dumps(sync_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )

    try:
        with urllib.request.urlopen(req) as resp:
            resp_data = resp.read().decode("utf-8")
            print(f"✅ Sync API Response: {resp_data}")
    except Exception as e:
        print(f"❌ Sync API Error: {e}")

if __name__ == "__main__":
    main()
