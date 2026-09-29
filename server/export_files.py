import os
import sys
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

_SERVER_DIR = os.path.dirname(os.path.abspath(__file__))
_OUTPUT_DIR = os.path.join(_SERVER_DIR, "outputs")
os.makedirs(_OUTPUT_DIR, exist_ok=True)

import build_dataset as b
records = b.unique_teachers

def export_dual_sheet_excel(records: list, output_filepath: str):
    df = pd.DataFrame(records)
    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    for col in columns:
        if col not in df.columns:
            df[col] = ''
            
    df = df.sort_values(by=['学校', '学院', '姓名']).reset_index(drop=True)
    df_with_email = df[df['邮箱'].fillna('').astype(str).str.strip() != ""].copy().reset_index(drop=True)
    df_no_email = df[df['邮箱'].fillna('').astype(str).str.strip() == ""].copy().reset_index(drop=True)
    
    df_with_email['序号'] = range(1, len(df_with_email) + 1)
    df_no_email['序号'] = range(1, len(df_no_email) + 1)
    
    df_with_email = df_with_email[columns]
    df_no_email = df_no_email[columns]
    
    with pd.ExcelWriter(output_filepath, engine='openpyxl') as writer:
        df_with_email.to_excel(writer, sheet_name='教师列表', index=False)
        df_no_email.to_excel(writer, sheet_name='无邮箱教师列表', index=False)
        
    wb = openpyxl.load_workbook(output_filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    content_font = Font(name="Microsoft YaHei", size=10, color="334155")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    col_widths = {
        'A': 8,   # 序号
        'B': 12,  # 姓名
        'C': 15,  # 职称
        'D': 25,  # 邮箱
        'E': 18,  # 学校
        'F': 18,  # 学院
        'G': 12,  # 采集人
        'H': 14,  # 采集日期
        'I': 45,  # 简介
        'J': 35   # 主页
    }
    
    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
            cell.border = thin_border
            
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = content_font
                cell.border = thin_border
                col_letter = get_column_letter(cell.column)
                if col_letter in ['A', 'B', 'C', 'G', 'H']:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                elif col_letter in ['D', 'J']:
                    cell.alignment = Alignment(horizontal='left', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                    
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width
            
    wb.save(output_filepath)
    print(f"Exported: {output_filepath}")

# Export both college excel and university excel
college_file = os.path.join(_OUTPUT_DIR, "甘肃中医药大学_护理学院.xlsx")
univ_file = os.path.join(_OUTPUT_DIR, "甘肃中医药大学.xlsx")

export_dual_sheet_excel(records, college_file)
export_dual_sheet_excel(records, univ_file)
print("Files created successfully.")
