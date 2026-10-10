"""Standard dual-sheet Excel export (有邮箱 / 无邮箱) shared by college and university deliveries.

Moved out of the retired ``autonomous_agent`` module so the delivery path no
longer imports the old script harvester (and PyMuPDF / Playwright with it).
"""
from __future__ import annotations

import os
import re
import tempfile
from zipfile import BadZipFile
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from server.delivery_fields import FIELD_ALIASES, STANDARD_FIELDS
from server import biography_workbook


EXCEL_COLUMNS = STANDARD_FIELDS
EXCEL_CELL_LIMIT = 32767
_INVALID_XML_TEXT = re.compile('[\ud800-\udfff\ufffe\uffff]')


def excel_cell_issues(records: list) -> list[dict]:
    """Detect cells Excel cannot represent without truncation; never edit records."""
    issues = []
    for index, record in enumerate(records):
        for column in EXCEL_COLUMNS:
            for key in (column, *FIELD_ALIASES.get(column, ())):
                value = record.get(key)
                if value is None:
                    continue
                text = _excel_safe_text(str(value))
                if _INVALID_XML_TEXT.search(text):
                    issues.append({'status': 'excel_invalid_xml_character', 'record_index': index,
                                   'field': column})
                    break
                if column != '简介' and biography_workbook.cell_units(text) > EXCEL_CELL_LIMIT:
                    issues.append({'status': 'excel_cell_too_long', 'record_index': index,
                                   'field': column, 'length': len(text), 'limit': EXCEL_CELL_LIMIT})
                    break
    return issues


def _excel_safe_text(value):
    """Drop characters an .xlsx file cannot hold; keep every visible character."""
    if isinstance(value, str):
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _store_formulas_as_text(workbook) -> None:
    """Harvested text is data: a value like ``=HYPERLINK(...)`` stays literal text."""
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def biography_storage_receipt(records: list) -> dict | None:
    """Describe lossless storage required for the normalized export records."""
    from server.delivery_fields import normalize_export_records
    safe = [{key: _excel_safe_text(value) for key, value in row.items()}
            for row in normalize_export_records(records)]
    _, _, receipts = biography_workbook.prepare(safe)
    if not receipts:
        return None
    return {'schema': 1, 'sheet': biography_workbook.SHEET, 'receipts': receipts}


def validate_biography_storage_receipt(records: list, receipt: dict | None) -> None:
    """The leased worker must explicitly attest storage of every long biography.

    This compares record-bound metadata; build_payload checks the actual workbook
    before issuing it, and university publishing verifies its own new artifact.
    """
    expected = biography_storage_receipt(records)
    if expected != receipt:
        raise ValueError('excel_cell_too_long: complete biography storage receipt missing or changed')


def verify_biography_storage(path: str | Path, records: list) -> dict | None:
    """Check complete biographies in the artifact against the native records."""
    receipt = biography_storage_receipt(records)
    if receipt is None:
        return None
    from server.delivery_fields import normalize_export_records
    safe = [{key: _excel_safe_text(value) for key, value in row.items()}
            for row in normalize_export_records(records)]
    try:
        workbook = openpyxl.load_workbook(path)
    except (BadZipFile, OSError) as exc:
        raise ValueError('full biography workbook unavailable or invalid') from exc
    try:
        biography_workbook.verify(workbook, safe)
    finally:
        workbook.close()
    return receipt


def export_dual_sheet_excel(records: list, output_filepath: str) -> dict:
    """Export styled dual-sheet standard Excel file adhering to strict invariants.

    Every value is harvested website text and is written as text: a value
    starting with ``=`` is never stored as a formula, and XML control
    characters (which openpyxl cannot store) are removed instead of failing
    the whole university workbook.
    """
    issues = excel_cell_issues(records)
    if issues:
        issue = issues[0]
        raise ValueError(f"{issue['status']}: record {issue['record_index']}, field {issue['field']}")
    from server.delivery_fields import normalize_export_records
    records = normalize_export_records(records)
    destination = Path(output_filepath)
    with tempfile.NamedTemporaryFile(dir=destination.parent, prefix=f'.{destination.stem}.',
                                     suffix='.xlsx', delete=False) as temporary:
        pending = Path(temporary.name)
    try:
        counts = _write_workbook(records, str(pending))
        os.replace(pending, destination)
        return counts
    finally:
        pending.unlink(missing_ok=True)


def _write_workbook(records: list, output_filepath: str) -> dict:
    """Serialize completely before the public export replaces its destination."""
    safe_records = [{key: _excel_safe_text(value) for key, value in row.items()} for row in records]
    copies, full_rows, receipts = biography_workbook.prepare(safe_records)
    df = pd.DataFrame(copies) if copies else pd.DataFrame(copies)
    columns = list(EXCEL_COLUMNS)
    if df.empty:
        empty_df = pd.DataFrame(columns=columns)
        with pd.ExcelWriter(output_filepath, engine='openpyxl') as writer:
            empty_df.to_excel(writer, sheet_name='教师列表', index=False)
            empty_df.to_excel(writer, sheet_name='无邮箱教师列表', index=False)
        wb = openpyxl.load_workbook(output_filepath)
        try:
            header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
            for ws in wb.worksheets:
                ws.views.sheetView[0].showGridLines = True
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = 16
        except Exception:
            pass
        wb.save(output_filepath)
        return {"total": 0, "with_email": 0, "no_email": 0}

    # Sort consistently without collapsing homonyms
    df = df.sort_values(by=['学校', '学院', '姓名']).reset_index(drop=True)

    df_with_email = df[df['邮箱'].fillna('').astype(str).str.strip() != ""].copy().reset_index(drop=True)
    df_no_email = df[df['邮箱'].fillna('').astype(str).str.strip() == ""].copy().reset_index(drop=True)

    df_with_email['序号'] = range(1, len(df_with_email) + 1)
    df_no_email['序号'] = range(1, len(df_no_email) + 1)

    for c in columns:
        if c not in df_with_email.columns:
            df_with_email[c] = ""
        if c not in df_no_email.columns:
            df_no_email[c] = ""

    links = [(name, index + 2, int(row['_biography_row']))
             for name, frame in [('教师列表', df_with_email), ('无邮箱教师列表', df_no_email)]
             for index, row in frame.iterrows()
             if '_biography_row' in row and pd.notna(row['_biography_row'])]
    df_with_email = df_with_email[columns]
    df_no_email = df_no_email[columns]

    with pd.ExcelWriter(output_filepath, engine='openpyxl') as writer:
        df_with_email.to_excel(writer, sheet_name='教师列表', index=False)
        df_no_email.to_excel(writer, sheet_name='无邮箱教师列表', index=False)
        if full_rows:
            pd.DataFrame(full_rows, columns=biography_workbook.HEADERS).to_excel(
                writer, sheet_name=biography_workbook.SHEET, index=False)

    wb = openpyxl.load_workbook(output_filepath)
    from openpyxl.worksheet.hyperlink import Hyperlink
    for name, row, first in links:
        cell = wb[name].cell(row, columns.index('简介') + 1)
        cell.hyperlink = Hyperlink(ref=cell.coordinate,
                                   location=f"'{biography_workbook.SHEET}'!A{first}", display=cell.value)
    _store_formulas_as_text(wb)  # must persist even if styling below fails

    # Style Excel Sheets with openpyxl
    try:
        header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
        data_font = Font(name="Microsoft YaHei", size=10, color="334155")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for ws in wb.worksheets:
            ws.views.sheetView[0].showGridLines = True
            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                for cell in row:
                    cell.border = thin_border
                    if row_idx == 1:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.font = data_font
                        if cell.column in [1, 2, 3, 5, 6, 7, 8]:
                            cell.alignment = Alignment(horizontal='center', vertical='center')
                        else:
                            cell.alignment = Alignment(horizontal='left', vertical='center')

            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                max_len = 0
                for cell in col[:15]:
                    val_str = str(cell.value or '')
                    max_len = max(max_len, len(val_str.encode('gbk', 'ignore')))
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)
    except Exception:
        pass
    wb.save(output_filepath)
    checked = openpyxl.load_workbook(output_filepath)
    try:
        biography_workbook.verify(checked, safe_records)
    finally:
        checked.close()

    return {
        "total": len(df),
        "with_email": len(df_with_email),
        "no_email": len(df_no_email),
        **({'full_biographies': len(receipts), 'biography_storage': {
            'schema': 1, 'sheet': biography_workbook.SHEET, 'receipts': receipts}} if receipts else {}),
    }
