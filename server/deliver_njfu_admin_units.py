import os
import sys

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import datetime
import json
import httpx

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

tasks = [
    {
        'task_id': '92b11b2e',
        'school': '南京林业大学',
        'college': '研究生院',
        'reason': '纯研究生培养/招生行政职能管理部门，无独立专任教师编制（导师均归属各专业学院）'
    },
    {
        'task_id': '5e3cf276',
        'school': '南京林业大学',
        'college': '继续教育学院',
        'reason': '纯学历继续教育与自考网络培训管理机构，无独立专任教师编制（教学任务由专业学院承担）'
    },
    {
        'task_id': '1f83e1b8',
        'school': '南京林业大学',
        'college': '国际教育学院',
        'reason': '承担中外合作办学与留学生管理协调职能部门，无独立专任教师编制（授课教师由专业学院承担）'
    }
]

output_dir = r"d:\Agent Workspace\Intern\MyPlan\server\outputs"
os.makedirs(output_dir, exist_ok=True)
date_str = datetime.datetime.now().strftime("%Y-%m-%d")
agent_name = "Antigravity (Playwright Engine)"

results = []

for t in tasks:
    timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{t['school']}_{t['college']}_{timestamp_str}.xlsx"
    filepath = os.path.join(output_dir, filename)
    
    # 0 records standard dual sheet excel
    columns = ['序号', '姓名', '职称', '邮箱', '学校', '学院', '采集人', '采集日期', '简介', '主页']
    df_empty = pd.DataFrame(columns=columns)
    
    sheet1_name = "有邮箱学者 (共 0 位)"
    sheet2_name = "无邮箱学者 (共 0 位)"
    
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df_empty.to_excel(writer, sheet_name=sheet1_name, index=False)
        df_empty.to_excel(writer, sheet_name=sheet2_name, index=False)
        
    # Style Excel
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E293B")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for row in ws.iter_rows(min_row=1, max_row=1):
            for cell in row:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center', vertical='center')
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = 16
    wb.save(filepath)
    
    sync_payload = {
        "records": [],
        "stats": {
            "total": 0,
            "with_email": 0,
            "no_email": 0,
            "email_rate": 0.0
        },
        "download_url": f"/download/{filename}",
        "progress": 100
    }
    
    res = httpx.post(f"http://127.0.0.1:8000/api/task/{t['task_id']}/sync", json=sync_payload, timeout=10)
    print(f"{t['task_id']} ({t['college']}) synced: status={res.status_code}, file={filename}")
    results.append({
        "task_id": t["task_id"],
        "school": t["school"],
        "college": t["college"],
        "total": 0,
        "with_email": 0,
        "no_email": 0,
        "excel": f"server/outputs/{filename}",
        "note": t["reason"]
    })

print("SUMMARY_JSON:" + json.dumps(results, ensure_ascii=False))
